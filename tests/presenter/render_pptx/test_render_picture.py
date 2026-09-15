"""Tests for image and plot-block rendering in the pptx engine."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from presenter.render_pptx import render, render_document


def _pictures(slide: Any) -> list[Any]:
    return [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]


def test_image_from_filesystem_path_honors_width_in(
    tmp_path: Path, png_image: Path
) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {
                "type": "image",
                "source": str(png_image),
                "width_in": 3.0,
                "caption": "Top-level schematic",
            }
        ],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    (picture,) = _pictures(slide)
    assert picture.width == Inches(3.0)
    captions = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text == "Top-level schematic"
    ]
    assert len(captions) == 1
    assert captions[0].top > picture.top


def test_plot_block_inserts_picture_from_source(
    tmp_path: Path, png_image: Path
) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [{"type": "plot", "source": str(png_image), "width_in": 2.5}],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    (picture,) = _pictures(slide)
    assert picture.width == Inches(2.5)


def test_image_without_width_uses_picture_placeholder(
    tmp_path: Path,
    png_image: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    render_document(
        {
            "title": "Deck",
            "slides": [
                {
                    "title": "Figure",
                    "layout": "Picture with Caption",
                    "blocks": [{"type": "image", "source": str(png_image)}],
                }
            ],
        },
        pptx_binding(),
        output,
    )
    slide = Presentation(str(output)).slides[-1]
    (picture,) = [
        shape
        for shape in slide.shapes
        if shape.is_placeholder and shape.placeholder_format.type.name == "PICTURE"
    ]
    assert picture.is_placeholder is True


def test_image_wider_than_slide_is_scaled_to_content_width(
    tmp_path: Path, png_image: Path
) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [{"type": "image", "source": str(png_image), "width_in": 99.0}],
        None,
        output,
    )
    prs = Presentation(str(output))
    slide = prs.slides[0]
    (picture,) = _pictures(slide)
    assert picture.width <= prs.slide_width - 2 * Inches(0.5)


def test_image_artifact_reference_resolves_via_binding(
    tmp_path: Path,
    png_image: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    binding = pptx_binding(artifacts={"artifact://schematics/ldo-top": str(png_image)})
    render(
        [
            {
                "type": "image",
                "source": "artifact://schematics/ldo-top",
                "width_in": 2.0,
            }
        ],
        binding,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    assert len(_pictures(slide)) == 1


def test_unresolvable_image_source_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="not found on disk"):
        render(
            [{"type": "image", "source": "/nonexistent/figure.png"}],
            None,
            tmp_path / "deck.pptx",
        )


def test_mapped_but_missing_artifact_raises(
    tmp_path: Path, pptx_binding: Callable[..., dict]
) -> None:
    binding = pptx_binding(artifacts={"artifact://missing": "/nonexistent/figure.png"})
    with pytest.raises(FileNotFoundError, match="artifact://missing"):
        render(
            [{"type": "image", "source": "artifact://missing", "width_in": 1.0}],
            binding,
            tmp_path / "deck.pptx",
        )


def test_image_block_without_source_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="requires a 'source' key"):
        render([{"type": "image"}], None, tmp_path / "deck.pptx")
