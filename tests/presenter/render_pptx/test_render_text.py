"""Tests for text-block rendering in the pptx engine."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from presenter.render_pptx import render, render_document


def _placeholder_by_name(slide: Any, name: str) -> Any:
    return {ph.name: ph for ph in slide.placeholders}[name]


def _text_boxes(slide: Any) -> list[Any]:
    return [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
    ]


def test_body_text_fills_body_placeholder(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render([{"type": "text", "style": "body", "text": "A 3.3 V LDO."}], None, output)
    slide = Presentation(str(output)).slides[0]
    body = _placeholder_by_name(slide, "Content Placeholder 2")
    assert body.text_frame.text == "A 3.3 V LDO."


def test_heading1_fills_title_placeholder(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render([{"type": "text", "style": "heading1", "text": "Overview"}], None, output)
    slide = Presentation(str(output)).slides[0]
    assert slide.shapes.title.text == "Overview"


def test_second_body_block_appends_paragraph(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {"type": "text", "style": "body", "text": "First point"},
            {"type": "text", "style": "body", "text": "Second point"},
        ],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    frame = _placeholder_by_name(slide, "Content Placeholder 2").text_frame
    assert [p.text for p in frame.paragraphs] == ["First point", "Second point"]


def test_heading2_and_caption_render_as_styled_text_boxes(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {"type": "text", "style": "heading2", "text": "Section"},
            {"type": "text", "style": "caption", "text": "Figure note"},
        ],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    boxes = {box.text_frame.text: box for box in _text_boxes(slide)}
    assert set(boxes) == {"Section", "Figure note"}
    heading_run = boxes["Section"].text_frame.paragraphs[0].runs[0]
    assert heading_run.font.bold is True
    assert heading_run.font.size == Pt(20.0)
    caption_run = boxes["Figure note"].text_frame.paragraphs[0].runs[0]
    assert caption_run.font.italic is True


def test_text_style_defaults_to_body(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render([{"type": "text", "text": "implicit body"}], None, output)
    slide = Presentation(str(output)).slides[0]
    body = _placeholder_by_name(slide, "Content Placeholder 2")
    assert body.text_frame.text == "implicit body"


def test_body_text_box_fallback_on_layout_without_body(
    tmp_path: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    render_document(
        {
            "title": "Deck",
            "slides": [
                {
                    "title": "Blank slide",
                    "layout": "Blank",
                    "blocks": [{"type": "text", "style": "body", "text": "free text"}],
                }
            ],
        },
        pptx_binding(),
        output,
    )
    slide = Presentation(str(output)).slides[0]
    boxes = [box for box in _text_boxes(slide) if box.text_frame.text == "free text"]
    assert len(boxes) == 1


def test_binding_layout_roles_drive_placeholder_choice(
    tmp_path: Path,
    make_template_pptx: Callable[..., Path],
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    template = make_template_pptx(rename_layouts={1: "Corp Content"})
    output = tmp_path / "deck.pptx"
    binding = pptx_binding(
        template=template,
        layouts={
            "Corp Content": {"Title 1": "title", "Content Placeholder 2": "title"}
        },
    )
    slide_spec = {
        "title": "Corp section",
        "layout": "Corp Content",
        "blocks": [{"type": "text", "style": "heading1", "text": "Override target"}],
    }
    render_document({"title": "Deck", "slides": [slide_spec]}, binding, output)
    slide = Presentation(str(output)).slides[-1]
    # The binding maps the content placeholder to the title role, so the
    # heading resolves there instead of the layout's own title placeholder.
    assert _placeholder_by_name(slide, "Content Placeholder 2").text_frame.text == (
        "Override target"
    )
    assert _placeholder_by_name(slide, "Title 1").text_frame.text == "Corp section"


def test_placeholder_role_falls_back_to_placeholder_type(
    tmp_path: Path,
    make_template_pptx: Callable[..., Path],
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    template = make_template_pptx()
    output = tmp_path / "deck.pptx"
    slide_spec = {
        "title": "Titled",
        "layout": "Title Only",
        "blocks": [{"type": "text", "style": "heading1", "text": "Heading"}],
    }
    render_document(
        {"title": "Deck", "slides": [slide_spec]},
        pptx_binding(template=template),
        output,
    )
    slide = Presentation(str(output)).slides[-1]
    # No binding entry for "Title Only": roles come from placeholder types,
    # so the slide title takes the layout's title placeholder and the
    # heading1 block, finding it already used, falls back to a text box.
    assert _placeholder_by_name(slide, "Title 1").text_frame.text == "Titled"
    boxes = [b for b in _text_boxes(slide) if b.text_frame.text == "Heading"]
    assert len(boxes) == 1


def test_theme_fonts_apply_to_free_text_boxes(
    tmp_path: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    binding = pptx_binding(
        styles={"theme": {"fonts": {"major": "Widget Head", "minor": "Widget Body"}}}
    )
    render(
        [{"type": "text", "style": "caption", "text": "themed caption"}],
        binding,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    box = next(b for b in _text_boxes(slide) if b.text_frame.text == "themed caption")
    run = box.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Widget Body"
