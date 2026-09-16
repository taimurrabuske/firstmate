"""Tests for native vector SVG embedding with asvg:svgBlip and PNG fallback."""

from __future__ import annotations

import io
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from presenter._ooxml.constants import (
    NS_A,
    NS_ASVG,
    NS_R,
    SVG_EXTENSION_URI,
)
from presenter.core.errors import MissingRendererError
from presenter.figure import Figure
from presenter.render_pptx import render, render_document

_CIRCUIT_SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200" width="400" height="200">'
    '<rect width="400" height="200" fill="#f8f9fa"/>'
    '<path d="M 50 100 L 120 100 L 140 70 L 160 130 L 180 70 L 200 130 L 220 100 L 350 100" '
    'stroke="#003366" stroke-width="3" fill="none"/>'
    '<circle cx="50" cy="100" r="5" fill="#cc0000"/>'
    '<circle cx="350" cy="100" r="5" fill="#009900"/>'
    '<text x="200" y="40" font-family="Arial" font-size="16" text-anchor="middle" fill="#333333">'
    'Precision Voltage Divider</text>'
    '</svg>'
)


@pytest.fixture()
def svg_file(tmp_path: Path) -> Path:
    """Create a sample vector SVG schematic."""
    p = tmp_path / "schematic.svg"
    p.write_text(_CIRCUIT_SVG, encoding="utf-8")
    return p


@pytest.fixture()
def png_fallback(tmp_path: Path) -> Path:
    """Create a sample PNG fallback."""
    p = tmp_path / "schematic_fallback.png"
    Image.new("RGB", (400, 200), color=(240, 240, 245)).save(p, format="PNG")
    return p


def _extract_slide_relationships(z: zipfile.ZipFile, slide_num: int = 1) -> dict[str, dict[str, str]]:
    """Parse slide rels XML into a dictionary of rId -> {Type, Target}."""
    rels_path = f"ppt/slides/_rels/slide{slide_num}.xml.rels"
    rels_xml = z.read(rels_path).decode("utf-8")
    root = ET.fromstring(rels_xml)
    result = {}
    for rel in root:
        r_id = rel.attrib.get("Id", "")
        r_type = rel.attrib.get("Type", "")
        target = rel.attrib.get("Target", "")
        result[r_id] = {"Type": r_type, "Target": target}
    return result


def test_svg_and_png_dual_relationship_packaging(
    tmp_path: Path, svg_file: Path, png_fallback: Path
) -> None:
    """Validate generated PPTX has valid SVG extension parts and relationship IDs."""
    output = tmp_path / "deck.pptx"
    render(
        [
            {
                "type": "image",
                "source": str(svg_file),
                "fallback": str(png_fallback),
                "width_in": 4.0,
                "caption": "Precision Voltage Divider Schematic",
                "alt_text": "Circuit schematic showing resistor divider",
            }
        ],
        None,
        output,
    )

    assert output.is_file()

    with zipfile.ZipFile(output) as z:
        namelist = z.namelist()
        svg_parts = [n for n in namelist if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in namelist if n.startswith("ppt/media/") and n.endswith(".png")]

        assert len(svg_parts) >= 1, f"Expected SVG part in media, found {namelist}"
        assert len(png_parts) >= 1, f"Expected PNG part in media, found {namelist}"
        assert z.read(svg_parts[0]) == _CIRCUIT_SVG.encode("utf-8")

        # 1. Inspect slide relationships
        rels = _extract_slide_relationships(z, slide_num=1)
        image_rels = {k: v for k, v in rels.items() if v["Type"].endswith("/image")}
        assert len(image_rels) >= 2, f"Expected at least 2 image relationships, got {image_rels}"

        png_target = Path(png_parts[0]).name
        svg_target = Path(svg_parts[0]).name

        png_r_id = next(k for k, v in image_rels.items() if png_target in v["Target"])
        svg_r_id = next(k for k, v in image_rels.items() if svg_target in v["Target"])
        assert png_r_id != svg_r_id, "PNG and SVG relationship IDs must be distinct"

        # 2. Inspect slide XML for DrawingML picture structure
        slide_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
        root = ET.fromstring(slide_xml)

        blip = root.find(f".//{{{NS_A}}}blip")
        assert blip is not None
        assert blip.attrib.get(f"{{{NS_R}}}embed") == png_r_id, (
            f"a:blip r:embed must point to PNG fallback {png_r_id}"
        )

        ext = blip.find(f".//{{{NS_A}}}ext[@uri='{SVG_EXTENSION_URI}']")
        assert ext is not None, f"Expected a:ext with URI {SVG_EXTENSION_URI}"

        svg_blip = ext.find(f"{{{NS_ASVG}}}svgBlip")
        assert svg_blip is not None, "Expected asvg:svgBlip element inside extension"
        assert svg_blip.attrib.get(f"{{{NS_R}}}embed") == svg_r_id, (
            f"asvg:svgBlip r:embed must point to SVG part {svg_r_id}"
        )

        # 3. Check alt text
        c_nv_pr = root.find(f".//{{{NS_A}}}picLocks/../../..//{{{NS_A}}}cNvPr")
        if c_nv_pr is None:
            # PresentationML picture non-visual properties
            for el in root.iter():
                if el.tag.endswith("cNvPr") and "descr" in el.attrib:
                    c_nv_pr = el
                    break
        assert c_nv_pr is not None
        assert c_nv_pr.attrib.get("descr") == "Circuit schematic showing resistor divider"

        # 4. Check [Content_Types].xml
        ct_xml = z.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="svg" ContentType="image/svg+xml"' in ct_xml

    # 5. Reopen with python-pptx to verify clean package parsing
    prs = Presentation(str(output))
    slide = prs.slides[0]
    pictures = [s for s in slide.shapes if s.has_text_frame is False and hasattr(s, "image")]
    assert len(pictures) == 1
    assert pictures[0].width == Inches(4.0)


def test_figure_embeds_both_representations(tmp_path: Path, svg_file: Path, png_fallback: Path) -> None:
    """Verify Figure with SVG and PNG representations embeds both cleanly."""
    fig = Figure(
        svg=svg_file,
        png=png_fallback,
        caption="Dual Representation Figure",
        alt_text="Schematic vector art with fallback",
        width_in=5.0,
    )

    output = tmp_path / "figure_deck.pptx"
    render([fig.to_block()], None, output)

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        assert len(svg_parts) == 1
        assert len(png_parts) == 1


def test_only_svg_provided_cleanly_rasterizes_png_fallback(tmp_path: Path, svg_file: Path) -> None:
    """Verify when only SVG is provided, a high-res PNG fallback is cleanly rasterized."""
    output = tmp_path / "svg_only_deck.pptx"
    render(
        [
            {
                "type": "image",
                "source": str(svg_file),
                "width_in": 3.5,
                "caption": "Auto-rasterized fallback test",
            }
        ],
        None,
        output,
    )

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        assert len(svg_parts) == 1, "Expected SVG part to be embedded"
        assert len(png_parts) == 1, "Expected auto-generated PNG fallback to be embedded"

        # Verify rasterized PNG content
        png_data = z.read(png_parts[0])
        assert png_data.startswith(b"\x89PNG")
        with Image.open(io.BytesIO(png_data)) as img:
            assert img.format == "PNG"
            assert img.width > 0
            assert img.height > 0

    # Ensure Presentation reopens cleanly
    prs = Presentation(str(output))
    assert len(prs.slides) == 1


def test_extensionless_svg_source_without_xml_prolog_renders_as_vector(
    tmp_path: Path,
) -> None:
    """Verify a plain 'source' image block still embeds as SVG, not raster.

    Regression test: an SVG asset with no ".svg" filename extension and no
    leading "<?xml" declaration (e.g. a license/generator comment followed
    by the "<svg" tag past the content-sniffing window) must still be
    detected as SVG. Previously this misdetection let raw SVG bytes reach
    ``add_picture``, which raised ``PIL.UnidentifiedImageError`` because
    python-pptx opens embedded image blobs with Pillow to sniff their format.
    """
    long_comment = "<!-- " + ("Generated by ACME EDA Schematic Capture. " * 10) + "-->\n"
    svg_body = (
        long_comment
        + '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100" '
        'width="100" height="100"><rect width="100" height="100" fill="red"/></svg>'
    )
    assert svg_body.find("<svg") > 256, "fixture must exceed the old sniff window"

    source = tmp_path / "schematic_asset"  # no ".svg" suffix, as from an artifact store
    source.write_text(svg_body, encoding="utf-8")

    output = tmp_path / "extensionless_svg_deck.pptx"
    render([{"type": "image", "source": str(source)}], None, output)

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        assert len(svg_parts) == 1, "Expected SVG part to be embedded"
        assert len(png_parts) == 1, "Expected auto-generated PNG fallback to be embedded"

    # Ensure Presentation and every embedded image reopen cleanly (the
    # PIL.UnidentifiedImageError this guards against surfaces here).
    prs = Presentation(str(output))
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "image"):
                assert shape.image.blob


def test_only_svg_missing_renderer_raises(tmp_path: Path, svg_file: Path) -> None:
    """Verify MissingRendererError when only SVG is provided and no rasterizer is found."""
    output = tmp_path / "should_fail.pptx"
    with (
        patch("shutil.which", return_value=None),
        patch("presenter._ooxml.rasterize._rasterize_with_cairosvg", return_value=False),
        pytest.raises(MissingRendererError, match="Rasterizing SVG to PNG fallback requires"),
    ):
        render(
            [{"type": "image", "source": str(svg_file)}],
            None,
            output,
        )


def test_caller_supplied_fallback_is_honored(
    tmp_path: Path, svg_file: Path, png_fallback: Path
) -> None:
    """Verify explicit caller-supplied fallback bypasses rasterization."""
    output = tmp_path / "caller_fallback.pptx"
    # Even if rasterizers were absent, an explicit fallback must succeed
    with patch("shutil.which", return_value=None):
        render(
            [
                {
                    "type": "image",
                    "source": str(svg_file),
                    "fallback": str(png_fallback),
                }
            ],
            None,
            output,
        )

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        assert len(svg_parts) == 1
        assert len(png_parts) == 1


def test_media_deduplication_across_slides(
    tmp_path: Path, svg_file: Path, png_fallback: Path
) -> None:
    """Verify inserting the same SVG on multiple slides produces only one media part."""
    output = tmp_path / "dedup.pptx"
    render(
        [
            {"type": "image", "source": str(svg_file), "fallback": str(png_fallback)},
            {"type": "pagebreak"},
            {"type": "image", "source": str(svg_file), "fallback": str(png_fallback)},
        ],
        None,
        output,
    )

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        png_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        assert len(svg_parts) == 1, f"Expected 1 deduplicated SVG part, got {svg_parts}"
        assert len(png_parts) == 1, f"Expected 1 deduplicated PNG part, got {png_parts}"

    prs = Presentation(str(output))
    assert len(prs.slides) == 2


def test_svg_picture_placeholder_insertion(
    tmp_path: Path,
    svg_file: Path,
    png_fallback: Path,
    pptx_binding: Any,
) -> None:
    """Verify SVG picture insertion into a layout with a PICTURE placeholder."""
    output = tmp_path / "placeholder_svg.pptx"
    render_document(
        {
            "title": "Deck with SVG Placeholder",
            "slides": [
                {
                    "title": "Schematic Slide",
                    "layout": "Picture with Caption",
                    "blocks": [
                        {
                            "type": "image",
                            "source": str(svg_file),
                            "fallback": str(png_fallback),
                        }
                    ],
                }
            ],
        },
        pptx_binding(),
        output,
    )

    with zipfile.ZipFile(output) as z:
        slide_parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        slide_xml = z.read(slide_parts[-1]).decode("utf-8")
        assert SVG_EXTENSION_URI in slide_xml
        assert "svgBlip" in slide_xml


def test_artifact_reference_resolution_with_svg(
    tmp_path: Path,
    svg_file: Path,
    png_fallback: Path,
    pptx_binding: Any,
) -> None:
    """Verify SVG and fallback paths resolve through binding['artifacts']."""
    output = tmp_path / "artifact_svg.pptx"
    binding = pptx_binding(
        artifacts={
            "artifact://schematics/vdiv-svg": str(svg_file),
            "artifact://schematics/vdiv-png": str(png_fallback),
        }
    )
    render(
        [
            {
                "type": "image",
                "source": "artifact://schematics/vdiv-svg",
                "fallback": "artifact://schematics/vdiv-png",
                "width_in": 3.0,
            }
        ],
        binding,
        output,
    )

    with zipfile.ZipFile(output) as z:
        svg_parts = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".svg")]
        assert len(svg_parts) == 1


def test_equation_svg_image_rendering(
    tmp_path: Path, svg_file: Path, png_fallback: Path
) -> None:
    """Verify equation block carrying an SVG image embeds dual relationships."""
    output = tmp_path / "equation_svg.pptx"
    render(
        [
            {
                "type": "equation",
                "latex": r"\frac{V_{out}}{V_{in}} = \frac{R_2}{R_1 + R_2}",
                "image": str(svg_file),
                "fallback": str(png_fallback),
            }
        ],
        None,
        output,
    )

    with zipfile.ZipFile(output) as z:
        slide_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
        assert SVG_EXTENSION_URI in slide_xml
        assert "svgBlip" in slide_xml


@pytest.mark.skipif(
    not shutil.which("libreoffice") and not shutil.which("soffice"),
    reason="LibreOffice not available for headless validation",
)
def test_libreoffice_headless_clean_conversion(
    tmp_path: Path, svg_file: Path, png_fallback: Path
) -> None:
    """Ensure generated PPTX with asvg:svgBlip opens and converts without repair warnings."""
    output = tmp_path / "clean_open.pptx"
    render(
        [
            {
                "type": "image",
                "source": str(svg_file),
                "fallback": str(png_fallback),
                "width_in": 5.0,
                "caption": "Validation Schematic",
            }
        ],
        None,
        output,
    )

    tool = shutil.which("libreoffice") or shutil.which("soffice")
    assert tool is not None
    cmd = [
        tool,
        "--headless",
        "--convert-to",
        "pdf",
        str(output),
        "--outdir",
        str(tmp_path),
    ]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=60, check=False)
    assert res.returncode == 0, f"LibreOffice conversion failed: {res.stderr}"
    pdf_path = tmp_path / "clean_open.pdf"
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 0
