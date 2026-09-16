"""Tests for native OMML equation rendering in PPTX."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from presenter.omml import EquationConversionError, UnsupportedMacroError
from presenter.render_pptx import render

OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"


def test_render_transfer_function_native_pptx(tmp_path: Path) -> None:
    """Verify 2nd-order transfer function renders as native DrawingML text math (a14:m) in PPTX."""
    output = tmp_path / "transfer_function.pptx"
    latex = r"H(s) = \frac{K \cdot \omega_0^2}{s^2 + \frac{\omega_0}{Q} s + \omega_0^2}"

    render(
        [
            {
                "type": "equation",
                "latex": latex,
                "mode": "native",
                "caption": "Transfer function H(s)",
            }
        ],
        None,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    # Check text box containing the math container
    text_boxes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert len(text_boxes) >= 1

    tb = text_boxes[0]
    p = tb.text_frame.paragraphs[0]
    a14_m_elems = p._p.findall(f"{{{A14_NS}}}m")
    assert len(a14_m_elems) == 1

    omath = a14_m_elems[0].find(f"{{{OMML_NS}}}oMath")
    assert omath is not None

    # Check for fraction element <m:f>
    fracs = omath.findall(f".//{{{OMML_NS}}}f")
    assert len(fracs) >= 1

    # Check caption box
    assert len(text_boxes) >= 2
    assert "Transfer function H(s)" in text_boxes[1].text_frame.text


def test_render_differential_equation_native_pptx(tmp_path: Path) -> None:
    """Verify differential equation renders within a14:m in PPTX."""
    output = tmp_path / "diff_eq.pptx"
    latex = r"\frac{d^2 y}{dt^2} + 2\zeta\omega_n \frac{dy}{dt} + \omega_n^2 y = f(t)"

    render(
        [{"type": "equation", "latex": latex, "mode": "native"}],
        None,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    tb = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
    p = tb.text_frame.paragraphs[0]
    a14_m = p._p.find(f"{{{A14_NS}}}m")
    assert a14_m is not None
    omath = a14_m.find(f"{{{OMML_NS}}}oMath")
    assert omath is not None
    assert omath.find(f".//{{{OMML_NS}}}f") is not None


def test_render_matrix_equation_native_pptx(tmp_path: Path) -> None:
    """Verify matrix equation renders within a14:m in PPTX."""
    output = tmp_path / "matrix_eq.pptx"
    latex = (
        r"\begin{bmatrix} \dot{x}_1 \\ \dot{x}_2 \end{bmatrix} = "
        r"\begin{bmatrix} 0 & 1 \\ -\omega_n^2 & -2\zeta\omega_n \end{bmatrix} "
        r"\begin{bmatrix} x_1 \\ x_2 \end{bmatrix} + \begin{bmatrix} 0 \\ 1 \end{bmatrix} u"
    )

    render(
        [{"type": "equation", "latex": latex, "mode": "native"}],
        None,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    tb = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
    p = tb.text_frame.paragraphs[0]
    a14_m = p._p.find(f"{{{A14_NS}}}m")
    assert a14_m is not None
    omath = a14_m.find(f"{{{OMML_NS}}}oMath")
    assert omath is not None

    # Check matrices and accents
    matrices = omath.findall(f".//{{{OMML_NS}}}m")
    assert len(matrices) >= 4
    accents = omath.findall(f".//{{{OMML_NS}}}acc")
    assert len(accents) >= 2


def test_render_binding_equation_mode_native_pptx(tmp_path: Path) -> None:
    """Verify binding equation_mode='native' enables native math in PPTX."""
    output = tmp_path / "binding_native.pptx"
    binding = {"equation_mode": "native"}

    render(
        [
            {"type": "equation", "latex": r"E = mc^2"},
            {"type": "equation", "latex": r"V_{out} = A_v \cdot V_{in}"},
        ],
        binding,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    boxes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert len(boxes) == 2
    for box in boxes:
        p = box.text_frame.paragraphs[0]
        assert p._p.find(f"{{{A14_NS}}}m") is not None


def test_render_native_preferred_fallback_to_image_pptx(
    tmp_path: Path, png_image: Path
) -> None:
    """Verify unsupported complex macro falls back to image rendering in PPTX."""
    output = tmp_path / "fallback.pptx"
    latex = r"\unknowncomplexmacro{x}"

    render(
        [
            {
                "type": "equation",
                "latex": latex,
                "mode": "native-preferred",
                "image": str(png_image),
            }
        ],
        None,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1


def test_render_native_preferred_dynamic_image_fallback_pptx(tmp_path: Path) -> None:
    """Verify unsupported OMML construct falls back to dynamic matplotlib image rendering in PPTX."""
    output = tmp_path / "fallback_dynamic.pptx"
    latex = r"\boldsymbol{W} + \boldsymbol{X}"

    render(
        [{"type": "equation", "latex": latex, "mode": "native-preferred"}],
        None,
        output,
    )

    prs = Presentation(str(output))
    slide = prs.slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1


def test_render_native_required_raises_on_unsupported_macro_pptx(
    tmp_path: Path,
) -> None:
    """Verify native-required mode raises when unsupported macro is encountered in PPTX."""
    output = tmp_path / "error.pptx"
    latex = r"\unsupportedcomplexmacro{x}"

    with pytest.raises((UnsupportedMacroError, EquationConversionError)):
        render(
            [{"type": "equation", "latex": latex, "mode": "native-required"}],
            None,
            output,
        )
