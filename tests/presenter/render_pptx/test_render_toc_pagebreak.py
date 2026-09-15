"""Tests for toc and pagebreak slide-flow behavior in the pptx engine."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

from pptx import Presentation

from presenter.render_pptx import render, render_document


def _slide_titles(prs: Any) -> list[str | None]:
    return [
        slide.shapes.title.text if slide.shapes.title is not None else None
        for slide in prs.slides
    ]


def test_pagebreak_starts_new_slide(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {"type": "text", "style": "body", "text": "before"},
            {"type": "pagebreak"},
            {"type": "text", "style": "body", "text": "after"},
        ],
        None,
        output,
    )
    prs = Presentation(str(output))
    assert len(prs.slides) == 2


def test_toc_on_content_slide_starts_divider_slide(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {"type": "text", "style": "body", "text": "content first"},
            {"type": "toc"},
        ],
        None,
        output,
    )
    prs = Presentation(str(output))
    assert len(prs.slides) == 2


def test_toc_on_untouched_slide_stays_that_slide(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {"type": "toc"},
            {"type": "text", "style": "body", "text": "on the divider slide"},
        ],
        None,
        output,
    )
    prs = Presentation(str(output))
    assert len(prs.slides) == 1


def test_document_slide_with_toc_stays_one_slide(
    tmp_path: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    render_document(
        {
            "title": "Deck",
            "slides": [
                {"title": "Contents", "blocks": [{"type": "toc"}]},
                {
                    "title": "Overview",
                    "blocks": [{"type": "text", "style": "body", "text": "body"}],
                },
            ],
        },
        pptx_binding(),
        output,
    )
    prs = Presentation(str(output))
    assert _slide_titles(prs) == ["Contents", "Overview"]


def test_pagebreak_inside_document_slide_adds_slide(
    tmp_path: Path,
    pptx_binding: Callable[..., dict[str, Any]],
) -> None:
    output = tmp_path / "deck.pptx"
    render_document(
        {
            "title": "Deck",
            "slides": [
                {
                    "title": "One",
                    "blocks": [
                        {"type": "pagebreak"},
                        {"type": "text", "style": "body", "text": "continued"},
                    ],
                }
            ],
        },
        pptx_binding(),
        output,
    )
    prs = Presentation(str(output))
    assert len(prs.slides) == 2
    titles = _slide_titles(prs)
    assert titles[0] == "One"
    # The pagebreak slide inherits the layout's empty title placeholder.
    assert not titles[1]
