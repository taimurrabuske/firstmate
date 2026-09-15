"""Shared fixtures for presenter.render_pptx tests.

Builds every .pptx and image fixture in-test with python-pptx and Pillow;
no binary fixtures are committed.

Importing ``presenter.render_pptx`` never mutates the shared engine
registry (its ``register()`` is an explicit opt-in), so module-level
imports are safe here. Only tests that exercise the registry request the
``pptx_engine`` fixture, which registers the lane's engine for one test
and restores the built-in stub in teardown.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Callable

import pytest
from PIL import Image
from pptx import Presentation

REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


@pytest.fixture()
def png_image(tmp_path: Path) -> Path:
    """A tiny real PNG so python-pptx can embed it."""
    path = tmp_path / "figure.png"
    Image.new("RGB", (40, 20), color=(20, 90, 200)).save(path, format="PNG")
    return path


@pytest.fixture()
def make_template_pptx(tmp_path: Path) -> Callable[..., Path]:
    """Build a minimal corporate .pptx template carrying a marker slide."""

    def _make(
        name: str = "template.pptx",
        marker: str = "Corporate Master",
        rename_layouts: dict[int, str] | None = None,
    ) -> Path:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        slide.shapes.title.text = marker
        for index, layout_name in (rename_layouts or {}).items():
            prs.slide_layouts[index].name = layout_name
        path = tmp_path / name
        prs.save(str(path))
        return path

    return _make


@pytest.fixture()
def pptx_binding() -> Callable[..., dict[str, Any]]:
    """Build a TemplateBinding dict per the shared contract."""

    def _make(
        template: Path | None = None,
        layouts: dict[str, dict[str, str]] | None = None,
        placeholders: dict[str, list[dict[str, Any]]] | None = None,
        styles: dict[str, Any] | None = None,
        masters: list[dict[str, Any]] | None = None,
        artifacts: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        binding: dict[str, Any] = {
            "format": "pptx",
            "template": str(template) if template else "",
            "name": "Test Binding",
            "layouts": dict(layouts or {}),
            "placeholders": dict(placeholders or {}),
            "styles": dict(styles or {}),
        }
        if masters:
            binding["masters"] = list(masters)
        if artifacts:
            binding["artifacts"] = dict(artifacts)
        return binding

    return _make


@pytest.fixture()
def pptx_engine() -> Any:
    """Register the real pptx engine in core's registry for one test."""
    import presenter.render_pptx
    from presenter.core.render import unregister_engine

    presenter.render_pptx.register()
    yield presenter.render_pptx
    unregister_engine("pptx")
