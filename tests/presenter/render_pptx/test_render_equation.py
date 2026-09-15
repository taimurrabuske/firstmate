"""Tests for equation-block rendering in the pptx engine."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from presenter.render_pptx import render


def test_equation_with_pre_rendered_image_inserts_picture(
    tmp_path: Path, png_image: Path
) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {
                "type": "equation",
                "latex": "V_{out}",
                "image": str(png_image),
                "font_size_pt": 12,
            }
        ],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    pictures = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1


def test_equation_without_image_falls_back_to_latex_text_box(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render(
        [
            {
                "type": "equation",
                "latex": r"V_{out} = V_{ref} (1 + R_1 / R_2)",
                "font_size_pt": 12,
            }
        ],
        None,
        output,
    )
    slide = Presentation(str(output)).slides[0]
    boxes = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        and shape.text_frame.text == r"V_{out} = V_{ref} (1 + R_1 / R_2)"
    ]
    assert len(boxes) == 1
    run = boxes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Courier New"
    assert run.font.size == Pt(12.0)


def test_equation_text_box_defaults_to_18_pt(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render([{"type": "equation", "latex": "E = mc^2"}], None, output)
    slide = Presentation(str(output)).slides[0]
    (box,) = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        and shape.text_frame.text == "E = mc^2"
    ]
    assert box.text_frame.paragraphs[0].runs[0].font.size == Pt(18.0)


def test_equation_image_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        render(
            [{"type": "equation", "latex": "x", "image": "/nonexistent/eq.png"}],
            None,
            tmp_path / "deck.pptx",
        )
