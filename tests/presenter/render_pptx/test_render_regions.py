"""Saved-output regressions for native region ownership and helper integration."""

from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from presenter.core import DocumentSpec, SlideSpec
from presenter.core.errors import SpecError
from presenter.core.render import render as core_render, unregister_engine
from presenter.layouts import (
    bullets_and_table, image_grid, split_half_text_images, title_bullets,
    title_single_image, two_column_bullets,
)
from presenter.render_pptx import render, render_document


def _tables(slide):
    return [shape for shape in slide.shapes if shape.has_table]


def _text(slide, text):
    return next(s for s in slide.shapes if s.has_text_frame and s.text == text)


def _disjoint(a, b):
    return (a.left + a.width <= b.left or b.left + b.width <= a.left
            or a.top + a.height <= b.top or b.top + b.height <= a.top)


def _table_template(tmp_path, height=4):
    prs = Presentation()
    layout = prs.slide_layouts[1]
    ph = layout.placeholders[1]
    ph._element.xpath(".//p:ph")[0].set("type", "tbl")
    ph.left, ph.top, ph.width, ph.height = Inches(2), Inches(2), Inches(5), Inches(height)
    path = tmp_path / "table-template.pptx"
    prs.save(path)
    return path, layout.name


@pytest.mark.parametrize("caption", [None, "A native table caption"])
def test_native_table_placeholder_reopens_as_frame_and_continues(tmp_path, caption):
    template, layout = _table_template(tmp_path)
    output = tmp_path / "native.pptx"
    rows = [[str(i), f"value {i}"] for i in range(27)]
    render_document({"slides": [{"title": "Native", "layout": layout, "blocks": [{
        "type": "table", "headers": ["A", "B"], "rows": rows,
        "caption": caption, "alignments": ["left", "right"],
    }]}]}, {"template": str(template)}, output)
    prs = Presentation(output)
    recovered = []
    widths = []
    assert len(prs.slides) > 1
    for i, slide in enumerate(prs.slides):
        frame, = _tables(slide)
        assert frame.is_placeholder
        assert frame.placeholder_format.idx == 1
        assert frame.left == Inches(2)
        assert frame.width == Inches(5)
        assert frame.top >= Inches(2)
        assert frame.top + frame.height <= Inches(6)
        assert frame.height == sum(row.height for row in frame.table.rows)
        widths.append([col.width for col in frame.table.columns])
        assert [c.text for c in frame.table.rows[0].cells] == ["A", "B"]
        recovered.extend([[c.text for c in row.cells] for row in list(frame.table.rows)[1:]])
        assert slide.shapes.title.text == ("Native" if i == 0 else "Native (continued)")
        if caption and i == 0:
            box = _text(slide, caption)
            assert box.left == frame.left
            assert box.top + box.height <= frame.top
        elif caption:
            assert not any(s.has_text_frame and s.text == caption for s in slide.shapes)
    assert recovered == rows
    assert all(w == widths[0] for w in widths)


@pytest.mark.parametrize("reverse", [False, True])
def test_body_and_table_never_share_occupied_region(tmp_path, reverse):
    blocks = [{"type": "text", "text": "Owned body"},
              {"type": "table", "headers": ["A"], "rows": [["one"]], "caption": "Caption"}]
    if reverse:
        blocks.reverse()
    output = tmp_path / "owned.pptx"
    render_document({"slides": [{"title": "Mixed", "blocks": blocks}]}, None, output)
    prs = Presentation(output)
    texts = []
    for slide in prs.slides:
        texts.extend(s.text for s in slide.shapes if s.has_text_frame)
        for table in _tables(slide):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text in ("Owned body", "Caption"):
                    assert _disjoint(shape, table)
            assert table.top + table.height <= prs.slide_height - Inches(.5)
    assert texts.count("Owned body") == 1
    assert texts.count("Caption") == 1
    if not reverse:
        assert len(prs.slides) == 2
        assert not _tables(prs.slides[0])
        assert _text(prs.slides[1], "Caption")


@pytest.mark.parametrize("position", [0, 15])
def test_unsplittable_data_row_rejected_without_saving(tmp_path, position):
    output = tmp_path / "oversize.pptx"
    rows = [[str(i)] for i in range(position)] + [["line\n" * 100]]
    with pytest.raises(ValueError, match="cannot fit"):
        render([{"type": "table", "headers": ["Header"], "rows": rows}], None, output)
    assert not output.exists()


def test_unsplittable_header_rejected(tmp_path):
    with pytest.raises(ValueError, match="cannot fit"):
        render([{"type": "table", "headers": ["line\n" * 100], "rows": []}],
               None, tmp_path / "header.pptx")


def test_header_only_uses_exact_native_budget(tmp_path):
    template, layout = _table_template(tmp_path, height=.4)
    output = tmp_path / "header-only.pptx"
    render_document({"slides": [{"layout": layout, "blocks": [
        {"type": "table", "headers": ["Header"], "rows": []},
    ]}]}, {"template": str(template)}, output)
    prs = Presentation(output)
    assert len(prs.slides) == 1
    frame, = _tables(prs.slides[0])
    assert len(frame.table.rows) == 1
    assert frame.height == Inches(.4)


@pytest.mark.parametrize("caption", [None, "Empty table note"])
def test_empty_table_is_no_shape_with_optional_caption(tmp_path, caption):
    output = tmp_path / "empty.pptx"
    render([{"type": "table", "headers": [], "rows": [], "caption": caption}], None, output)
    prs = Presentation(output)
    assert len(prs.slides) == 1
    assert not _tables(prs.slides[0])
    if caption:
        assert _text(prs.slides[0], caption)


def _helper_cases(image):
    return [
        title_bullets("Overview", ["A", "B"]),
        title_single_image("Figure", image, caption="Figure caption"),
        split_half_text_images("Split", ["A", "B"], [image, image]),
        split_half_text_images("Split left", ["A", "B"], [image], side="left"),
        two_column_bullets("Columns", ["L1", "L2"], ["R1", "R2"]),
        two_column_bullets("Headed", ["L1"], ["R1"], left_heading="Left", right_heading="Right"),
        bullets_and_table("Table right", ["A"], ["H"], [["V"]], table_caption="Data"),
        bullets_and_table("Table left", ["A"], ["H"], [["V"]], side="left"),
        image_grid("Grid hint only", [image], columns=3),
    ]


@pytest.mark.parametrize("case", range(9))
def test_helpers_through_core_render_and_reopen(tmp_path, png_image, pptx_engine, case):
    helper = _helper_cases(str(png_image))[case]
    doc = DocumentSpec.from_dict({"title": "Helpers", "slides": [helper]})
    assert doc.to_dict()["slides"][0] == helper
    assert DocumentSpec.from_json(doc.to_json()).to_dict() == doc.to_dict()
    template = tmp_path / "template.pptx"
    Presentation().save(template)
    output = tmp_path / "helper.pptx"
    core_render(doc, {"format": "pptx", "template": str(template)}, output)
    prs = Presentation(output)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    if case == 4:
        left, right = _text(slide, "L1\nL2"), _text(slide, "R1\nR2")
        assert left.is_placeholder and right.is_placeholder
        assert _disjoint(left, right) and left.left < right.left
    if case == 5:
        left, right = _text(slide, "L1"), _text(slide, "R1")
        for heading, body in ((_text(slide, "Left"), left), (_text(slide, "Right"), right)):
            assert heading.is_placeholder and body.is_placeholder
            assert heading.top + heading.height <= body.top
        assert _disjoint(left, right)
    if case in (2, 3):
        body = _text(slide, "A\nB")
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        assert len(pictures) == (2 if case == 2 else 1)
        for picture in pictures:
            assert _disjoint(picture, body)
            assert picture.top >= body.top
            assert picture.top + picture.height <= body.top + body.height
            assert (picture.left > body.left) == (case == 2)
        if case == 2:
            assert _disjoint(*pictures)
    if case in (6, 7):
        table, = _tables(slide)
        body = _text(slide, "A")
        assert _disjoint(table, body)
        assert (table.left > body.left) == (case == 6)
        if case == 6:
            caption = _text(slide, "Data")
            assert caption.left == table.left
            assert _disjoint(caption, body) and _disjoint(caption, table)

    # The same normalized helper keeps its vertical reading order in Word.
    from presenter.render_docx import register
    template_docx = tmp_path / "template.docx"
    Document().save(template_docx)
    register()
    try:
        docx_output = tmp_path / "helper.docx"
        core_render(doc, {"format": "docx", "template": str(template_docx)}, docx_output)
    finally:
        unregister_engine("docx")
    word = Document(docx_output)
    paragraphs = [p.text for p in word.paragraphs if p.text]
    expected = [helper["title"]] + [b["text"] for b in helper["blocks"] if b["type"] == "text"]
    assert [p for p in paragraphs if p in expected] == expected
    assert len(word.tables) == sum(b["type"] == "table" for b in helper["blocks"])
    assert len(word.inline_shapes) == sum(b["type"] == "image" for b in helper["blocks"])


@pytest.mark.parametrize("side", ["left", "right"])
def test_peer_table_continuations_stay_in_native_region(tmp_path, pptx_engine, side):
    helper = bullets_and_table("Paged", ["Context"], ["H"], [[str(i)] for i in range(35)],
                               side=side, table_caption="Only once")
    template = tmp_path / "template.pptx"
    Presentation().save(template)
    output = tmp_path / "paged.pptx"
    core_render({"title": "Deck", "slides": [helper]},
                {"format": "pptx", "template": str(template)}, output)
    prs = Presentation(output)
    assert len(prs.slides) > 1
    assert _text(prs.slides[0], "Context").is_placeholder
    assert _text(prs.slides[0], "Only once")
    recovered = []
    for slide in prs.slides:
        table, = _tables(slide)
        assert (table.left > prs.slide_width // 2) == (side == "right")
        assert table.top + table.height <= prs.slide_height - Inches(.5)
        recovered.extend(row.cells[0].text for row in list(table.table.rows)[1:])
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text == "Context":
                assert _disjoint(shape, table)
    assert recovered == [str(i) for i in range(35)]


@pytest.mark.parametrize("extra", [
    {"placeholder_roles": {"left": [True]}},
    {"placeholder_roles": {"left": [-1]}},
    {"placeholder_roles": {"left": [1]}},
    {"placeholder_roles": {"left": [0], "right": [0]}},
    {"placeholder_roles": {"left": "0"}},
    {"placeholder_roles": None},
    {"placeholder_roles": {1: [0]}},
    {"columns": True}, {"columns": 0}, {"columns": "2"},
])
def test_invalid_helper_metadata_is_rejected(extra):
    with pytest.raises(SpecError):
        SlideSpec.from_dict({"blocks": [{"type": "text", "text": "A"}], **extra})


def test_helper_metadata_has_no_mutable_aliases():
    helper = two_column_bullets("Columns", ["A"], ["B"])
    spec = SlideSpec.from_dict(helper)
    helper["placeholder_roles"]["left"].clear()
    normalized = spec.to_dict()
    normalized["placeholder_roles"]["right"].clear()
    assert spec.placeholder_roles == {"left": [0], "right": [1]}


@pytest.mark.parametrize("mode", ["missing", "overlap"])
def test_peer_helpers_reject_incompatible_layouts(tmp_path, mode):
    helper = two_column_bullets("Columns", ["A"], ["B"])
    prs = Presentation()
    if mode == "missing":
        helper["layout"] = "Title and Content"
    else:
        peers = list(prs.slide_layouts[3].placeholders)
        peers[2].left = peers[1].left
    template = tmp_path / "bad-peers.pptx"
    prs.save(template)
    with pytest.raises(ValueError, match="native content placeholders|non-overlapping"):
        render_document({"slides": [helper]}, {"template": str(template)}, tmp_path / "out.pptx")


def test_multiple_large_peer_images_fit_without_overlap(tmp_path):
    from PIL import Image
    image = tmp_path / "tall.png"
    Image.new("RGB", (500, 1600)).save(image)
    helper = split_half_text_images("Split", ["Context"], [str(image), str(image)])
    for block in helper["blocks"][1:]:
        block["caption"] = "Caption " * 15
    output = tmp_path / "images.pptx"
    render_document({"slides": [helper]}, None, output)
    prs = Presentation(output)
    slide = prs.slides[0]
    content = [s for s in slide.shapes if (s.has_text_frame and s.text and s != slide.shapes.title)
               or s.shape_type == 13]
    for i, a in enumerate(content):
        assert a.top + a.height <= prs.slide_height - Inches(.5)
        for b in content[i + 1:]:
            assert _disjoint(a, b)


def test_empty_table_caption_moves_off_occupied_body(tmp_path):
    output = tmp_path / "empty-caption.pptx"
    render([{"type": "text", "text": "Body"},
            {"type": "table", "headers": [], "rows": [], "caption": "Note"}], None, output)
    prs = Presentation(output)
    assert len(prs.slides) == 2
    assert _text(prs.slides[0], "Body")
    assert _text(prs.slides[1], "Note")
    assert not any(_tables(slide) for slide in prs.slides)


def test_caption_cannot_force_a_row_off_slide(tmp_path):
    output = tmp_path / "too-tall.pptx"
    output.write_bytes(b"previous output must survive rejection")
    with pytest.raises(ValueError, match="cannot fit"):
        render([{"type": "table", "headers": ["H"], "rows": [["V"]],
                 "caption": "long caption " * 1000}], None, output)
    assert output.read_bytes() == b"previous output must survive rejection"
