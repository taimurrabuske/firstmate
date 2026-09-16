"""Tests for table-block rendering in the pptx engine."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from presenter.blocks.tables import (
    PVTCorner,
    SpecItem,
    build_pvt_table,
    build_spec_table,
)
from presenter.render_pptx import render, render_document


def _table_shapes(slide: Any) -> list[Any]:
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE]


def _render_table(block: dict, tmp_path: Path) -> Any:
    output = tmp_path / "deck.pptx"
    render([block], None, output)
    return Presentation(str(output)).slides[0]


def test_table_grid_renders_headers_and_rows(tmp_path: Path) -> None:
    slide = _render_table(
        {
            "type": "table",
            "headers": ["Parameter", "Min", "Typ", "Unit"],
            "rows": [
                ["Output voltage", 3.234, 3.3, "V"],
                ["Dropout", None, 150, "mV"],
            ],
            "style": "grid",
        },
        tmp_path,
    )
    (table,) = [shape.table for shape in _table_shapes(slide)]
    assert (len(table.rows), len(table.columns)) == (3, 4)
    assert table.cell(0, 0).text == "Parameter"
    assert table.cell(0, 3).text == "Unit"
    assert table.cell(1, 1).text == "3.234"
    assert table.cell(2, 1).text == ""  # null values render as empty cells
    assert table.cell(2, 2).text == "150"
    header_run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert header_run.font.bold is True
    data_run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert data_run.font.bold is not True


def test_table_grid_style_sets_bold_first_row_without_banding(tmp_path: Path) -> None:
    slide = _render_table(
        {"type": "table", "headers": ["A"], "rows": [["x"]], "style": "grid"},
        tmp_path,
    )
    (table,) = [shape.table for shape in _table_shapes(slide)]
    assert table.first_row is True
    assert table.horz_banding is False


def test_table_striped_style_sets_banding(tmp_path: Path) -> None:
    slide = _render_table(
        {"type": "table", "headers": ["A"], "rows": [["x"]], "style": "striped"},
        tmp_path,
    )
    (table,) = [shape.table for shape in _table_shapes(slide)]
    assert table.first_row is True
    assert table.horz_banding is True


def test_table_plain_style_clears_first_row_and_banding(tmp_path: Path) -> None:
    slide = _render_table(
        {"type": "table", "headers": ["A"], "rows": [["x"]], "style": "plain"},
        tmp_path,
    )
    (table,) = [shape.table for shape in _table_shapes(slide)]
    assert table.first_row is False
    assert table.horz_banding is False


def test_table_caption_sits_above_the_table(tmp_path: Path) -> None:
    slide = _render_table(
        {
            "type": "table",
            "caption": "Over PVT unless noted",
            "headers": ["A", "B"],
            "rows": [["x", "y"]],
        },
        tmp_path,
    )
    (graphic_frame,) = _table_shapes(slide)
    captions = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text == "Over PVT unless noted"
    ]
    assert len(captions) == 1
    assert captions[0].top < graphic_frame.top


def test_table_row_wider_than_headers_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="more than"):
        render(
            [
                {
                    "type": "table",
                    "headers": ["A", "B"],
                    "rows": [["x", "y", "z"]],
                }
            ],
            None,
            tmp_path / "deck.pptx",
        )


def test_table_without_headers_uses_row_shape(tmp_path: Path) -> None:
    slide = _render_table({"type": "table", "rows": [["a", "b"], ["c", "d"]]}, tmp_path)
    (table,) = [shape.table for shape in _table_shapes(slide)]
    assert (len(table.rows), len(table.columns)) == (2, 2)
    assert table.cell(0, 1).text == "b"


def test_large_table_paginates_across_multiple_slides(tmp_path: Path) -> None:
    output = tmp_path / "paginated.pptx"
    total_rows = 25
    headers = ["Parameter", "Min", "Typ", "Max"]
    rows = [
        [f"Param_{i}", f"{1.0 + i * 0.05:.2f}", f"{1.1 + i * 0.05:.2f}", f"{1.2 + i * 0.05:.2f}"]
        for i in range(total_rows)
    ]
    alignments = ["left", "right", "right", "right"]
    doc = {
        "title": "Deck",
        "slides": [
            {
                "title": "PLL Specifications",
                "blocks": [
                    {
                        "type": "table",
                        "headers": headers,
                        "rows": rows,
                        "alignments": alignments,
                        "style": "grid",
                    }
                ],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))

    # Must split across at least 2 slides (typically 3 slides for 25 rows)
    assert len(prs.slides) >= 2

    # Slide titles: first slide has original title, subsequent slides have (continued)
    assert prs.slides[0].shapes.title.text == "PLL Specifications"
    for slide in list(prs.slides)[1:]:
        assert slide.shapes.title.text == "PLL Specifications (continued)"

    # Check each slide has a table and valid layout (does not overflow slide height)
    data_rows_collected: list[list[str]] = []
    first_col_widths: list[int] | None = None

    for slide_idx, slide in enumerate(prs.slides):
        tbl_shapes = _table_shapes(slide)
        assert len(tbl_shapes) == 1, f"Slide {slide_idx} should have exactly 1 table"
        tbl_shape = tbl_shapes[0]
        table = tbl_shape.table

        # Layout boundary assertion: table must not exceed slide height
        assert tbl_shape.top + tbl_shape.height <= prs.slide_height

        # Column headers repeated on every slide
        assert [c.text for c in table.rows[0].cells] == headers
        for c_idx in range(len(headers)):
            cell = table.cell(0, c_idx)
            # Bold headers
            assert cell.text_frame.paragraphs[0].runs[0].font.bold is True
            # Header alignments match
            expected_align = PP_ALIGN.LEFT if alignments[c_idx] == "left" else PP_ALIGN.RIGHT
            assert cell.text_frame.paragraphs[0].alignment == expected_align

        # Preserved column widths across all slides
        col_widths = [col.width for col in table.columns]
        if first_col_widths is None:
            first_col_widths = col_widths
        else:
            assert col_widths == first_col_widths

        # Collect data rows and verify data alignments
        for r_idx in range(1, len(table.rows)):
            row_cells = [table.cell(r_idx, c).text for c in range(len(headers))]
            data_rows_collected.append(row_cells)
            for c_idx in range(len(headers)):
                cell = table.cell(r_idx, c_idx)
                expected_align = PP_ALIGN.LEFT if alignments[c_idx] == "left" else PP_ALIGN.RIGHT
                assert cell.text_frame.paragraphs[0].alignment == expected_align

    # Verify complete row distribution without loss or duplication
    assert len(data_rows_collected) == total_rows
    assert data_rows_collected == rows


def test_large_table_with_caption_places_caption_on_first_slide_only(tmp_path: Path) -> None:
    output = tmp_path / "caption_pagination.pptx"
    total_rows = 25
    headers = ["Corner", "Frequency"]
    rows = [[f"Corner_{i}", f"{100 + i * 10} MHz"] for i in range(total_rows)]
    caption_text = "Operating conditions across PVT"

    doc = {
        "title": "Deck",
        "slides": [
            {
                "title": "VCO Corners",
                "blocks": [
                    {
                        "type": "table",
                        "caption": caption_text,
                        "headers": headers,
                        "rows": rows,
                        "style": "striped",
                    }
                ],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))

    assert len(prs.slides) >= 2

    # First slide has the caption above the table
    first_slide_captions = [
        s for s in prs.slides[0].shapes
        if s.has_text_frame and s.text_frame.text == caption_text
    ]
    assert len(first_slide_captions) == 1
    (first_tbl,) = _table_shapes(prs.slides[0])
    assert first_slide_captions[0].top < first_tbl.top

    # Continuation slides must NOT repeat the caption text box
    for slide in list(prs.slides)[1:]:
        cont_captions = [
            s for s in slide.shapes
            if s.has_text_frame and s.text_frame.text == caption_text
        ]
        assert len(cont_captions) == 0


def test_spec_table_integration_pagination(tmp_path: Path) -> None:
    output = tmp_path / "spec_table.pptx"
    items = [
        SpecItem(
            f"Param_{i}",
            symbol=f"P{i}",
            min_val=1.0 + i * 0.01,
            typ_val=1.1 + i * 0.01,
            max_val=1.2 + i * 0.01,
            unit="V",
        )
        for i in range(25)
    ]
    spec_block = build_spec_table(items, caption="Table 1: Operating Limits")
    doc = {
        "title": "PLL Review",
        "slides": [
            {
                "title": "Specification Summary",
                "blocks": [spec_block],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))

    assert len(prs.slides) >= 2
    assert prs.slides[0].shapes.title.text == "Specification Summary"
    assert prs.slides[1].shapes.title.text == "Specification Summary (continued)"

    total_data_rows = 0
    for slide in prs.slides:
        (tbl_shape,) = _table_shapes(slide)
        table = tbl_shape.table
        # Header is repeated
        assert table.cell(0, 0).text == "Parameter"
        assert table.cell(0, 5).text == "Units"
        total_data_rows += len(table.rows) - 1

    assert total_data_rows == 25


def test_pvt_corner_table_integration_pagination(tmp_path: Path) -> None:
    output = tmp_path / "pvt_corners.pptx"
    corners = [
        PVTCorner(
            name=f"C{i:02d}",
            process="typical" if i % 2 == 0 else "fast",
            voltage=1.8 + i * 0.01,
            temperature=-40 + i * 5,
        )
        for i in range(30)
    ]
    pvt_block = build_pvt_table(corners, caption="Table 2: PVT Corners")
    doc = {
        "title": "PVT Review",
        "slides": [
            {
                "title": "Corner Analysis",
                "blocks": [pvt_block],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))

    assert len(prs.slides) >= 2
    assert prs.slides[0].shapes.title.text == "Corner Analysis"
    assert prs.slides[1].shapes.title.text == "Corner Analysis (continued)"

    total_data_rows = 0
    for slide in prs.slides:
        (tbl_shape,) = _table_shapes(slide)
        table = tbl_shape.table
        assert table.cell(0, 0).text == "Corner"
        assert table.cell(0, 1).text == "Process"
        total_data_rows += len(table.rows) - 1

    assert total_data_rows == 30


def test_table_without_headers_paginates_cleanly(tmp_path: Path) -> None:
    output = tmp_path / "no_headers_paginated.pptx"
    rows = [[f"cell_{r}_{c}" for c in range(3)] for r in range(25)]
    render([{"type": "table", "rows": rows}], None, output)
    prs = Presentation(str(output))

    assert len(prs.slides) >= 2
    total_data_rows = 0
    for slide in prs.slides:
        (tbl_shape,) = _table_shapes(slide)
        total_data_rows += len(tbl_shape.table.rows)

    assert total_data_rows == 25


def test_small_table_does_not_paginate(tmp_path: Path) -> None:
    output = tmp_path / "small_table.pptx"
    doc = {
        "title": "Deck",
        "slides": [
            {
                "title": "Small Table",
                "blocks": [
                    {
                        "type": "table",
                        "headers": ["A", "B"],
                        "rows": [["1", "2"], ["3", "4"]],
                    }
                ],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))
    assert len(prs.slides) == 1
    assert prs.slides[0].shapes.title.text == "Small Table"


def test_continuation_title_when_slide_has_no_initial_title(tmp_path: Path) -> None:
    output = tmp_path / "no_title_pagination.pptx"
    render(
        [
            {
                "type": "table",
                "headers": ["Col1", "Col2"],
                "rows": [[f"r{i}c1", f"r{i}c2"] for i in range(25)],
            }
        ],
        None,
        output,
    )
    prs = Presentation(str(output))
    assert len(prs.slides) >= 2
    for slide in list(prs.slides)[1:]:
        assert slide.shapes.title.text == "(continued)"


def test_slide_with_prior_content_splits_appropriately(tmp_path: Path) -> None:
    output = tmp_path / "prior_content.pptx"
    doc = {
        "title": "Deck",
        "slides": [
            {
                "title": "Mixed Slide",
                "blocks": [
                    {"type": "text", "style": "heading2", "text": "Introductory Section"},
                    {
                        "type": "table",
                        "headers": ["Col1", "Col2"],
                        "rows": [[f"r{i}c1", f"r{i}c2"] for i in range(20)],
                        "style": "grid",
                    },
                ],
            }
        ],
    }
    render_document(doc, None, output)
    prs = Presentation(str(output))
    assert len(prs.slides) >= 2
    assert prs.slides[0].shapes.title.text == "Mixed Slide"
    assert prs.slides[1].shapes.title.text == "Mixed Slide (continued)"
