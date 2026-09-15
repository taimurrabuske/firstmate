"""Tests for template-binding behavior of the pptx engine."""

from __future__ import annotations

from pathlib import Path
from typing import Callable

import pytest
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from presenter.render_pptx import render


def test_template_binding_starts_from_template_file(
    tmp_path: Path,
    make_template_pptx: Callable[..., Path],
    pptx_binding: Callable[..., dict],
) -> None:
    template = make_template_pptx(marker="Corporate Master")
    output = tmp_path / "out.pptx"
    result = render(
        [{"type": "text", "style": "body", "text": "Appended body"}],
        pptx_binding(template=template),
        output,
    )
    assert result["template"] == str(template)

    prs = Presentation(str(output))
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title.text == "Corporate Master"
    bodies = [
        ph.text_frame.text
        for ph in prs.slides[1].placeholders
        if ph.text_frame.text == "Appended body"
    ]
    assert len(bodies) == 1


def test_empty_template_field_renders_default_presentation(
    tmp_path: Path, pptx_binding: Callable[..., dict]
) -> None:
    output = tmp_path / "out.pptx"
    result = render(
        [{"type": "text", "style": "body", "text": "hello"}],
        pptx_binding(template=None),
        output,
    )
    assert result["template"] is None
    prs = Presentation(str(output))
    assert len(prs.slides) == 1


def test_missing_template_path_raises_instead_of_blank_fallback(
    tmp_path: Path, pptx_binding: Callable[..., dict]
) -> None:
    binding = pptx_binding(template=tmp_path / "no-such-template.pptx")
    with pytest.raises(PackageNotFoundError):
        render([{"type": "toc"}], binding, tmp_path / "x.pptx")


def test_non_pptx_binding_format_is_rejected(
    tmp_path: Path, pptx_binding: Callable[..., dict]
) -> None:
    binding = pptx_binding()
    binding["format"] = "docx"
    with pytest.raises(ValueError, match="format 'pptx', got 'docx'"):
        render([{"type": "toc"}], binding, tmp_path / "x.pptx")


def test_render_summary_reports_counts(tmp_path: Path) -> None:
    output = tmp_path / "out.pptx"
    result = render(
        [
            {"type": "text", "style": "body", "text": "one"},
            {"type": "pagebreak"},
            {"type": "toc"},
        ],
        None,
        output,
    )
    assert result == {
        "format": "pptx",
        "output_path": str(output),
        "template": None,
        "blocks_rendered": 3,
    }
    assert Presentation(str(output)) is not None


def test_render_creates_missing_output_parent_directory(tmp_path: Path) -> None:
    output = tmp_path / "nested" / "dir" / "out.pptx"
    render([{"type": "toc"}], None, output)
    assert output.is_file()


def test_empty_block_list_renders_single_slide(tmp_path: Path) -> None:
    output = tmp_path / "out.pptx"
    result = render([], None, output)
    assert result["blocks_rendered"] == 0
    assert len(Presentation(str(output)).slides) == 1


def test_unsupported_block_type_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="unsupported block type"):
        render([{"type": "mesh"}], None, tmp_path / "out.pptx")
