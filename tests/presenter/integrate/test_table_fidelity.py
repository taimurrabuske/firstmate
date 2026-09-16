"""Cross-engine regression coverage for the table null-cell and alignment
convention: presenter.render_docx.engine and presenter.render_pptx.engine
must treat the same table block identically for null-cell text and
per-column alignment, per the contract in presenter/README.md.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from presenter.render_docx import render as render_docx
from presenter.render_pptx import render as render_pptx

_ALIGNMENTS = ["left", "right", "center", "right"]

_TABLE_BLOCK = {
    "type": "table",
    "headers": ["Parameter", "Min", "Typ", "Unit"],
    "rows": [
        ["Output voltage", None, 3.3, "V"],
        ["Dropout", 100, None, "mV"],
    ],
    "style": "grid",
    "alignments": _ALIGNMENTS,
}

_DOCX_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}
_PPTX_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def test_docx_table_normalizes_null_and_applies_column_alignment(tmp_path: Path) -> None:
    output = tmp_path / "out.docx"
    render_docx([_TABLE_BLOCK], None, output)
    table = Document(str(output)).tables[0]

    assert table.cell(1, 1).text == ""  # null "Min" for Output voltage
    assert table.cell(2, 2).text == ""  # null "Typ" for Dropout

    for col, align in enumerate(_ALIGNMENTS):
        expected = _DOCX_ALIGN[align]
        assert table.cell(0, col).paragraphs[0].alignment == expected
        assert table.cell(1, col).paragraphs[0].alignment == expected
        assert table.cell(2, col).paragraphs[0].alignment == expected


def test_pptx_table_normalizes_null_and_applies_column_alignment(tmp_path: Path) -> None:
    output = tmp_path / "out.pptx"
    render_pptx([_TABLE_BLOCK], None, output)
    slide = Presentation(str(output)).slides[0]
    (table,) = [
        shape.table for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE
    ]

    assert table.cell(1, 1).text == ""  # null "Min" for Output voltage
    assert table.cell(2, 2).text == ""  # null "Typ" for Dropout

    for col, align in enumerate(_ALIGNMENTS):
        expected = _PPTX_ALIGN[align]
        assert table.cell(0, col).text_frame.paragraphs[0].alignment == expected
        assert table.cell(1, col).text_frame.paragraphs[0].alignment == expected
        assert table.cell(2, col).text_frame.paragraphs[0].alignment == expected
