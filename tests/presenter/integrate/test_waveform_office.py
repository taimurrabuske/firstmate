"""Native art at physical size in real Office packages and LibreOffice output."""

from __future__ import annotations

import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from presenter.blocks import build_plot
from presenter.core import DocumentBuilder, DocumentSpec, render
from presenter.core.errors import WaveformFigureError
from presenter.waveform_origin import (
    extract_waveform_bundle,
    read_office_waveform_origins,
)


@pytest.fixture
def office_files(tmp_path):
    if not (3, 12) <= sys.version_info[:2] < (3, 14):
        pytest.skip("native producer requires Python 3.12-3.13")
    pytest.importorskip("waveforms_widgets.native_figure")
    import presenter.render_docx
    import presenter.render_pptx

    presenter.render_docx.register(replace=True)
    presenter.render_pptx.register(replace=True)
    x = np.linspace(0, 2, 200)
    fig = build_plot(
        x=x,
        y=np.sin(2 * np.pi * x),
        title="Native signal",
        x_label="Time",
        x_unit="s",
        y_label="Output",
        y_unit="V",
        figsize=(4, 3),
        dpi=144,
        assets_dir=tmp_path / "sources",
        as_figure=True,
        identities=[{"source": "source-42", "run": "run-3", "signal": "vout"}],
    )
    spec = DocumentSpec.from_json(
        DocumentBuilder("Plots").slide("Waveforms").add(fig).build().to_json()
    )
    result = []
    for fmt, factory in (("pptx", Presentation), ("docx", Document)):
        template = tmp_path / f"template.{fmt}"
        factory().save(template)
        output = tmp_path / f"native-{fmt}.{fmt}"
        render(spec, {"format": fmt, "template": str(template)}, output)
        result.append((fmt, output))
    shutil.rmtree(tmp_path / "sources")
    return fig, result


def test_office_evidence_originals_and_ink_survive_source_removal(
    office_files, tmp_path
):
    from waveforms_widgets.native_figure import reopen_figure

    fig, files = office_files
    for fmt, output in files:
        origins = read_office_waveform_origins(output)
        assert origins == [fig.waveform_origin]
        _, waveforms = reopen_figure(
            extract_waveform_bundle(origins[0], tmp_path / f"reopened-{fmt}")
        )
        np.testing.assert_allclose(
            waveforms[0].y, np.sin(2 * np.pi * np.linspace(0, 2, 200))
        )
        with ZipFile(output) as archive:
            media = [
                n
                for n in archive.namelist()
                if n.startswith("ppt/media/" if fmt == "pptx" else "word/media/")
            ]
            assert any(n.endswith(".png") for n in media)
            if fmt == "pptx":
                assert any(n.endswith(".svg") for n in media)
                assert b"svgBlip" in archive.read("ppt/slides/slide1.xml")
            assert any(
                n.startswith("presenter/waveform-member") and n.endswith(".svg")
                for n in archive.namelist()
            )
        if fmt == "pptx":
            pictures = [
                s for s in Presentation(output).slides[0].shapes if s.shape_type == 13
            ]
            assert len(pictures) == 1
            width, height = pictures[0].width, pictures[0].height
        else:
            picture = Document(output).inline_shapes[0]
            width, height = picture.width, picture.height
        assert width == 4 * 914400
        assert height == 3 * 914400


@pytest.mark.parametrize("mutation", ["changed", "missing", "relationship"])
def test_office_reopen_rejects_bad_parts(office_files, tmp_path, mutation):
    _, files = office_files
    for fmt, output in files:
        bad = tmp_path / f"bad-{mutation}.{fmt}"
        with ZipFile(output) as source, ZipFile(bad, "w", ZIP_DEFLATED) as dest:
            victim = next(
                n
                for n in source.namelist()
                if n.startswith("presenter/waveform-member") and n.endswith(".npy")
            )
            for name in source.namelist():
                data = source.read(name)
                if name == victim:
                    if mutation == "missing":
                        continue
                    if mutation == "changed":
                        data = b"bad"
                if mutation == "relationship" and name.startswith("presenter/_rels/"):
                    root = ET.fromstring(data)
                    root.remove(root[0])
                    data = ET.tostring(root)
                dest.writestr(name, data)
        with pytest.raises(WaveformFigureError):
            read_office_waveform_origins(bad)


def test_actual_office_render_has_signal_labels_and_physical_size(
    office_files, tmp_path
):
    tools = [shutil.which(n) for n in ("libreoffice", "pdftoppm", "pdftotext")]
    if not all(tools):
        pytest.skip("actual Office output requires libreoffice and poppler")
    fig, files = office_files
    import base64
    import io

    from scipy.ndimage import label

    def signal_mask(pixels):
        # Largest connected blue component excludes the legend and Word heading.
        raw = (
            (pixels[:, :, 2].astype(int) - pixels[:, :, 0] > 65)
            & (pixels[:, :, 1] > 60)
            & (pixels[:, :, 1] < 180)
            & (pixels[:, :, 0] < 100)
        )
        components, _ = label(raw, structure=np.ones((3, 3)))
        counts = np.bincount(components.ravel())
        counts[0] = 0
        return components == counts.argmax()

    original = np.array(
        Image.open(
            io.BytesIO(base64.b64decode(fig.waveform_origin["members"]["figure.png"]))
        ).convert("RGB")
    )
    oy, ox = np.where(signal_mask(original))
    for fmt, output in files:
        profile = tmp_path / f"lo-{fmt}"
        result = subprocess.run(
            [
                tools[0],
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmp_path),
                str(output),
            ],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        assert result.returncode == 0, result.stderr
        pdf = output.with_suffix(".pdf")
        assert pdf.exists(), result.stdout + result.stderr
        raster = tmp_path / f"page-{fmt}"
        subprocess.run(
            [
                tools[1],
                "-r",
                "144",
                "-f",
                "1",
                "-singlefile",
                "-png",
                str(pdf),
                str(raster),
            ],
            check=True,
            timeout=60,
        )
        pixels = np.array(Image.open(raster.with_suffix(".png")).convert("RGB"))
        # Actual curve color survives the Office SVG/PDF path and Word PNG path.
        blue = signal_mask(pixels)
        ys, xs = np.where(blue)
        assert len(xs) > 500
        # Office must preserve native signal size at the requested 144 dpi,
        # independently of the producer's legend/axis layout allocation.
        assert 200 < xs.max() - xs.min() < 576
        assert 200 < ys.max() - ys.min() < 432
        assert abs((xs.max() - xs.min()) - (ox.max() - ox.min())) <= 3
        assert abs((ys.max() - ys.min()) - (oy.max() - oy.min())) <= 3
        assert np.count_nonzero(blue[:, xs.min() : xs.min() + 3]) > 3
        assert np.count_nonzero(blue[:, xs.max() - 2 : xs.max() + 1]) > 3
        # The vector presentation keeps real readable text; Word deliberately uses PNG.
        if fmt == "pptx":
            text = subprocess.run(
                [tools[2], str(pdf), "-"], capture_output=True, text=True, check=True
            ).stdout
            for text_label in ("Native signal", "Time", "Output", "[s]", "[V]"):
                assert text_label in text
