"""End-to-end test rendering a real .docx through the full presenter stack.

Exercises the same composition as examples/datasheet_demo.py: core spec
assembly, a fixture template inventoried by presenter.templates, an image
resolved through presenter.artifacts, technical content blocks from
presenter.blocks, and dispatch through presenter.core.render into the
presenter.render_docx engine.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from PIL import Image

import presenter.render_docx
from presenter.artifacts import FilesystemArtifactSource
from presenter.blocks.equation import build_equation
from presenter.blocks.plot import build_plot
from presenter.blocks.tables import SpecItem, build_spec_table
from presenter.core import DocumentBuilder, render
from presenter.core.blocks import image, pagebreak, text, toc
from presenter.templates.docx import inventory_docx_template


def _fixture_template(path: Path) -> None:
    document = Document()
    document.add_heading("Corporate Template Marker", level=1)
    document.save(str(path))


def _fixture_artifact(root: Path, ref: str) -> None:
    path = root / ref
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (16, 16), color=(10, 20, 30)).save(path, format="PNG")


@pytest.fixture()
def rendered_docx(tmp_path: Path) -> Path:
    presenter.render_docx.register()

    template_path = tmp_path / "template.docx"
    _fixture_template(template_path)
    binding = inventory_docx_template(template_path)

    artifacts_root = tmp_path / "artifacts"
    ref = "schematics/top.png"
    _fixture_artifact(artifacts_root, ref)
    source = FilesystemArtifactSource(artifacts_root)
    binding["artifacts"] = {ref: str(source.resolve_path(ref))}

    plot_block = build_plot(
        backend="ordinary",
        x=[0.0, 1.0, 2.0], y=[1.0, 2.0, 1.5], caption="trace", output_path=tmp_path / "p.png"
    )
    equation_block = build_equation("a+b", output_path=tmp_path / "e.png")
    table_block = build_spec_table(
        [SpecItem("Vout", symbol="V", min_val=3.2, typ_val=3.3, max_val=3.4, unit="V")],
        caption="specs",
    )

    spec = (
        DocumentBuilder("Demo Datasheet", author="Test author")
        .slide("Contents").add(toc())
        .slide("Overview")
        .add(text("Overview text", style="body"))
        .add(image(ref, width_in=2.0, caption="schematic"))
        .slide("Specs")
        .add(table_block)
        .add(equation_block)
        .add(pagebreak())
        .slide("Waveform")
        .add(plot_block)
        .build()
    )

    output_path = tmp_path / "out.docx"
    written = render(spec, binding, output_path)
    assert written == output_path
    return output_path


def test_output_file_is_written(rendered_docx: Path) -> None:
    assert rendered_docx.is_file()
    assert rendered_docx.stat().st_size > 0


def test_template_content_is_preserved(rendered_docx: Path) -> None:
    document = Document(str(rendered_docx))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert "Corporate Template Marker" in headings


def test_slide_titles_become_headings_in_document_order(rendered_docx: Path) -> None:
    document = Document(str(rendered_docx))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    # The template's own heading comes first, then one per slide, in order.
    assert headings == ["Corporate Template Marker", "Contents", "Overview", "Specs", "Waveform"]


def test_table_and_pictures_render(rendered_docx: Path) -> None:
    document = Document(str(rendered_docx))
    assert len(document.tables) == 1
    assert document.tables[0].rows[0].cells[0].text == "Parameter"
    # schematic (artifact-resolved), equation image, and plot: three pictures.
    assert len(document.inline_shapes) == 3


def test_body_text_and_toc_field_present(rendered_docx: Path) -> None:
    document = Document(str(rendered_docx))
    body_texts = [p.text for p in document.paragraphs]
    assert "Overview text" in body_texts
    assert any("Table of contents" in t for t in body_texts)
