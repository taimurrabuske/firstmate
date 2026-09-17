"""Native circuit artifacts through the ordinary builder/core/Office engines."""

from __future__ import annotations

import builtins
import hashlib
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from presenter import CircuitFigureError, Figure
from presenter.blocks import circuit_figure
from presenter.core import DocumentBuilder, DocumentSpec, render
from presenter.core.blocks import ImageBlock, PlotBlock, image, validate_block
from presenter.core.errors import PresenterError

NS = "{http://www.w3.org/2000/svg}"


@pytest.fixture
def native():
    # Ordinary Presenter installs do not require the optional native producer.
    pytest.importorskip("schematic.core.circuit_figure")
    from schematic import SchematicBuilder
    from schematic.core.schematic_data import AnalysisData, FreeTextData

    b = SchematicBuilder()
    r = b.add_resistor("R1", at=(0, 0), value="10k", rotation=90)
    c = b.add_capacitor("C1", at=(110, 0), value="2p")
    b.wire(pins=[r.pin("1"), c.pin("1")], route="hv", net_name="out")
    for component in (r, c):
        b.label(component, "model", show_value=False)
    b.label(r, "refdes", side="top")
    b.label(r, "value", side="top", order=1)
    for wire in b.data.wires:
        wire.label_offset = (45, -15)
    b.data.free_texts.append(FreeTextData(-30, 75, "tau = RC = 20 ns"))
    b.data.analysis_items.append(AnalysisData(-30, 110, ["R1"], [["V", "out"]], []))
    return b


@pytest.fixture
def figure(native, tmp_path):
    return circuit_figure(
        native.data,
        symbols=native.library,
        assets_dir=tmp_path / "assets",
        width_in=4,
        request={"metadata": {"evidence": {"source": "synthetic RC", "tau_s": 2e-8}}},
        caption="RC circuit",
    )


def test_native_default_and_evidence_roundtrip(figure, native):
    assert isinstance(figure, Figure)
    payload = figure.circuit_origin["payload"]
    assert {c["symbol"] for c in payload["connectivity"][""]} == {
        "RZ_RESISTOR",
        "RZ_CAPACITOR",
    }
    nets = {
        c["instance_id"]: {n["net"] for n in c["nodes"]}
        for c in payload["connectivity"][""]
    }
    assert "out" in nets["R1"] & nets["C1"]
    assert payload["document"]["schematic"] == json.loads(
        json.dumps(native.data.to_dict())
    )
    assert "tau = RC" in Path(figure.svg).read_text()
    assert Figure(**figure.to_dict()).to_block() == figure.to_block()
    for accessor, kind in ((ImageBlock, "image"), (PlotBlock, "plot")):
        typed = accessor.from_dict(figure.to_block(kind))
        assert typed.circuit_origin == figure.circuit_origin
        assert (
            validate_block(typed.to_dict())["circuit_origin"] == figure.circuit_origin
        )
    spec = DocumentBuilder("RC").slide("Circuit").add(figure).build()
    loaded = DocumentSpec.from_json(spec.to_json())
    assert loaded.to_dict() == spec.to_dict()
    assert (
        json.loads(json.dumps(figure.to_dict()))["circuit_origin"]
        == figure.circuit_origin
    )
    assert Image.open(figure.png).width == 1200


def test_physical_font_sizes_do_not_grow_with_raster_dpi(native, tmp_path):
    from PIL import ImageChops, ImageStat

    low = circuit_figure(
        native.data, symbols=native.library, assets_dir=tmp_path, width_in=4, dpi=96
    )
    high = circuit_figure(
        native.data, symbols=native.library, assets_dir=tmp_path, width_in=4, dpi=300
    )

    def on_white(path):
        rgba = Image.open(path).convert("RGBA")
        white = Image.new("RGBA", rgba.size, "white")
        return Image.alpha_composite(white, rgba).convert("RGB")

    baseline = on_white(low.png)
    scaled = on_white(high.png).resize(baseline.size, Image.Resampling.LANCZOS)
    assert max(ImageStat.Stat(ImageChops.difference(baseline, scaled)).mean) < 5


def test_envelope_and_closed_symbols(native, tmp_path):
    from schematic.core.envelope_adapter import encode_schematic_envelope

    source = encode_schematic_envelope(
        native.data,
        normalize_origin=False,
        revision=9,
        metadata={"analytical_source": {"run": "rc-9", "tau_s": 2e-8}},
    )
    f = circuit_figure(
        source,
        symbols=native.library.bundle_text(
            [c.symbol_name for c in native.data.components]
        ),
        assets_dir=tmp_path,
    )
    assert f.circuit_origin["payload"]["document"]["schematic"] == json.loads(
        json.dumps(native.data.to_dict())
    )
    assert f.circuit_origin["source"] == json.loads(json.dumps(source))


def test_explicit_images_are_not_inferred_from_caption(tmp_path):
    path = tmp_path / "generic.png"
    Image.new("RGB", (10, 10), "white").save(path)
    for block in (
        image(str(path), caption="RC circuit"),
        Figure(png=path, caption="Circuit schematic").to_block(),
    ):
        assert "circuit_origin" not in validate_block(block)


def test_missing_integration_has_presenter_root_error(monkeypatch, tmp_path):
    real_import = builtins.__import__

    def unavailable(name, *args, **kwargs):
        if name.startswith("schematic"):
            raise ImportError("integration absent")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", unavailable)
    with pytest.raises(PresenterError, match="Install the delivered Schematic"):
        circuit_figure({}, assets_dir=tmp_path)


@pytest.mark.parametrize("field", ["source", "svg", "png", "fallback", "width_in"])
def test_wrong_binding_refuses(figure, field):
    block = figure.to_block()
    block[field] = 2 if field == "width_in" else "wrong.svg"
    with pytest.raises(CircuitFigureError, match="binding|width_in"):
        validate_block(block)


def test_asset_digest_and_physical_size_are_bound_to_evidence(figure):
    block = figure.to_block()
    block["width_in"] = block["circuit_origin"]["width_in"] = 8
    with pytest.raises(CircuitFigureError, match="binding digest mismatch"):
        validate_block(block)
    block = figure.to_block()
    block["circuit_origin"]["assets"]["png"]["sha256"] = "0" * 64
    with pytest.raises(CircuitFigureError, match="binding digest mismatch"):
        validate_block(block)


def test_evidence_tamper_refuses(figure):
    block = figure.to_block()
    block["circuit_origin"]["payload"]["request"]["metadata"]["evidence"]["tau_s"] = 4
    with pytest.raises(CircuitFigureError, match="digest mismatch"):
        validate_block(block)


@pytest.mark.parametrize("kind", ["svg", "png"])
@pytest.mark.parametrize("action", ["missing", "stale"])
def test_missing_or_stale_asset_fails_before_dispatch(figure, tmp_path, kind, action):
    path = Path(getattr(figure, kind))
    if action == "missing":
        path.unlink()
    else:
        path.write_bytes(b"wrong asset")
    from presenter.core.render import register_engine

    def forbidden(*args):
        pytest.fail("engine must not receive stale circuit assets")

    register_engine("docx", forbidden, replace=True)
    with pytest.raises(CircuitFigureError, match="asset|No such file"):
        render(
            DocumentBuilder("bad").add(figure).build(),
            {"format": "docx", "template": "unused"},
            tmp_path / "bad.docx",
        )


def test_missing_rasterizer_is_actionable(native, tmp_path, monkeypatch):
    from presenter.core.errors import MissingRendererError

    def unavailable(*args, **kwargs):
        raise MissingRendererError("renderer missing")

    monkeypatch.setattr("presenter.blocks.circuit.rasterize_svg_to_png", unavailable)
    with pytest.raises(CircuitFigureError, match="Install CairoSVG"):
        circuit_figure(native.data, symbols=native.library, assets_dir=tmp_path)


@pytest.mark.parametrize("replacement", [[], {"request": []}, None, float("nan")])
def test_malformed_evidence_is_a_presenter_error(figure, replacement):
    block = figure.to_block()
    block["circuit_origin"]["payload"] = replacement
    with pytest.raises(PresenterError):
        validate_block(block)


def test_invalid_native_inputs_never_fall_back(native, tmp_path):
    from schematic.core.schematic_data import LatexFragmentData

    with pytest.raises(CircuitFigureError, match="explicit symbol closure"):
        circuit_figure(native.data, assets_dir=tmp_path)
    with pytest.raises(CircuitFigureError, match="symbol|Symbol"):
        circuit_figure(
            native.data,
            symbols="<svg xmlns='http://www.w3.org/2000/svg'/>",
            assets_dir=tmp_path,
        )
    native.data.latex_fragments.append(LatexFragmentData(0, 0, "x", "", 50, 20))
    with pytest.raises(CircuitFigureError, match="latex_fragments"):
        circuit_figure(native.data, symbols=native.library, assets_dir=tmp_path)
    assert not list(tmp_path.glob("*.png"))


def test_missing_child_wrong_occurrence_and_overlay_refuse(native, tmp_path):
    for request, message in (
        ({"selected_path": ["absent"]}, "Missing hierarchy instance"),
        (
            {
                "overlays": [
                    {
                        "target": ["absent"],
                        "text": "gain",
                        "source": "run",
                        "key": "gain",
                    }
                ]
            },
            "Missing overlay target",
        ),
        ({"overlays": [{"latex": "x"}]}, "overlay requires"),
    ):
        with pytest.raises(CircuitFigureError, match=message):
            circuit_figure(
                native.data,
                symbols=native.library,
                request=request,
                assets_dir=tmp_path,
            )
    native.define_block("child", ["in", "out"])
    native.place("child", "X1", at=(250, 0))
    with pytest.raises(CircuitFigureError, match="Missing child schematic"):
        circuit_figure(native.data, symbols=native.library, assets_dir=tmp_path)


def test_mathml_and_native_svg_do_not_compete_for_default_namespace(native, tmp_path):
    from presenter.omml.latex_to_mathml import latex_to_mathml

    svg = native.export_circuit_figure()
    for _ in range(2):
        math = latex_to_mathml(r"\frac{a}{b}")
        assert math.startswith('<math xmlns="http://www.w3.org/1998/Math/MathML"')
        assert (
            ET.fromstring(math).find("{http://www.w3.org/1998/Math/MathML}mfrac")
            is not None
        )
        assert native.export_circuit_figure() == svg
        circuit_figure(svg, assets_dir=tmp_path)


def test_artifact_visual_tamper_and_missing_metadata_refuse(native, tmp_path):
    svg = native.export_circuit_figure()
    with pytest.raises(CircuitFigureError, match="differs from its native source"):
        circuit_figure(svg.replace("tau = RC", "wrong label", 1), assets_dir=tmp_path)
    with pytest.raises(CircuitFigureError, match="metadata"):
        circuit_figure(
            b'<svg xmlns="http://www.w3.org/2000/svg"/>', assets_dir=tmp_path
        )
    with pytest.raises(CircuitFigureError, match="physical size"):
        circuit_figure(svg, assets_dir=tmp_path, width_in=8)
    with pytest.raises(CircuitFigureError, match="already bind"):
        circuit_figure(svg, assets_dir=tmp_path, request={})


def _office_files(figure, tmp_path):
    import presenter.render_docx
    import presenter.render_pptx

    presenter.render_docx.register(replace=True)
    presenter.render_pptx.register(replace=True)
    spec = DocumentSpec.from_json(
        DocumentBuilder("Circuit example")
        .slide("Native circuit")
        .add(figure)
        .build()
        .to_json()
    )
    for fmt, factory in (("docx", Document), ("pptx", Presentation)):
        template = tmp_path / f"template.{fmt}"
        factory().save(template)
        out = tmp_path / f"circuit.{fmt}"
        render(spec, {"format": fmt, "template": str(template)}, out)
        yield fmt, out


def test_saved_office_contains_native_sources_evidence_and_correct_ink(
    figure, tmp_path
):
    for fmt, out in _office_files(figure, tmp_path):
        with ZipFile(out) as z:
            names = z.namelist()
            manifest_name = next(
                n
                for n in names
                if n.startswith("presenter/circuit-origin") and n.endswith(".json")
            )
            manifest = json.loads(z.read(manifest_name))
            item = manifest["figures"][0]
            assert item["circuit_origin"] == figure.circuit_origin
            for kind, part in item["package_parts"].items():
                content = z.read(part.lstrip("/"))
                assert (
                    hashlib.sha256(content).hexdigest()
                    == figure.circuit_origin["assets"][kind]["sha256"]
                )
            media = [
                n
                for n in names
                if n.startswith("ppt/media/" if fmt == "pptx" else "word/media/")
            ]
            assert any(n.endswith(".png") for n in media)
            if fmt == "pptx":
                assert any(n.endswith(".svg") for n in media)
                assert b"svgBlip" in z.read("ppt/slides/slide1.xml")
            else:
                assert not any(n.endswith(".svg") for n in media)
            assert any(
                b"urn:presenter:relationships:circuit-origin" in z.read(n)
                for n in names
                if n.endswith(".rels")
            )
        if fmt == "docx":
            assert Document(out).inline_shapes[0].width == 4 * 914400
        else:
            pictures = [
                s for s in Presentation(out).slides[0].shapes if s.shape_type == 13
            ]
            assert len(pictures) == 1
            assert pictures[0].width == 4 * 914400


def test_canvas_rc_and_two_level_hierarchy_artifacts(tmp_path):
    pytest.importorskip("canvas.circuit_figure")
    from canvas import module
    from canvas.app_services import state, symbol_library
    from canvas.circuit_figure import export_circuit_figure
    from schematic import SchematicBuilder

    module.build_application()
    design = state()
    design.reset()
    try:
        for name, source, selected in (
            ("rc", "R1 (in out) resistor r=10k\nC1 (out 0) capacitor c=2p\n", ()),
            (
                "hierarchy",
                (
                    "subckt leaf in out\nR1 (in out) resistor r=10k\nends leaf\n"
                    "subckt stage in out\nXleaf (in out) leaf\nends stage\nXstage (in out) stage\n"
                ),
                ("Xstage", "Xleaf"),
            ),
        ):
            design.reset()
            design.load_netlist(source, dut="top")
            if selected:
                builder = SchematicBuilder(library=symbol_library())
                builder.define_block("leaf", ["in", "out"])
                builder.define_block("stage", ["in", "out"])
            artifact = export_circuit_figure(
                design, capture=True, occurrence_path=selected, width_mm=101.6
            )
            f = circuit_figure(artifact.encode(), assets_dir=tmp_path / name)
            origin = f.circuit_origin
            assert origin["producer"] == "canvas"
            assert origin["payload"]["request"]["selected_path"] == list(selected)
            assert all(
                b["correspondence"]["state"] == "complete"
                for b in origin["payload"]["request"]["metadata"]["occurrences"]
            )
            if selected:
                assert set(origin["payload"]["request"]["children"]) == {
                    "stage",
                    "leaf",
                }
                tree = ET.fromstring(Path(f.svg).read_text())
                assert any(
                    g.get("data-instance-path") == '["Xstage","Xleaf","R1"]'
                    for g in tree.findall(f"{NS}g")
                )
            out_dir = tmp_path / name / "office"
            out_dir.mkdir()
            assert len(list(_office_files(f, out_dir))) == 2
    finally:
        design.reset()
