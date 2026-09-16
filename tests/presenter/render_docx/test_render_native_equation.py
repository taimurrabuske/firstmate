"""Tests for native OMML equation rendering in DOCX."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from presenter.omml import EquationConversionError, UnsupportedMacroError
from presenter.render_docx import render

OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def test_render_transfer_function_native_docx(
    tmp_path: Path, docx_binding: object
) -> None:
    """Verify 2nd-order transfer function renders as native m:oMathPara in DOCX."""
    output = tmp_path / "transfer_function.docx"
    latex = r"H(s) = \frac{K \cdot \omega_0^2}{s^2 + \frac{\omega_0}{Q} s + \omega_0^2}"

    render(
        [
            {
                "type": "equation",
                "latex": latex,
                "mode": "native",
                "caption": "Transfer function H(s)",
            }
        ],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    p = document.paragraphs[0]
    omath_paras = p._p.findall(f"{{{OMML_NS}}}oMathPara")
    assert len(omath_paras) == 1
    omath = omath_paras[0].find(f"{{{OMML_NS}}}oMath")
    assert omath is not None

    # Check for fraction element <m:f>
    fracs = omath.findall(f".//{{{OMML_NS}}}f")
    assert len(fracs) >= 1

    # Check caption paragraph below
    assert len(document.paragraphs) >= 2
    assert "Transfer function H(s)" in document.paragraphs[1].text


def test_render_differential_equation_native_docx(
    tmp_path: Path, docx_binding: object
) -> None:
    """Verify 2nd-order differential equation renders as native OMML in DOCX."""
    output = tmp_path / "diff_eq.docx"
    latex = r"\frac{d^2 y}{dt^2} + 2\zeta\omega_n \frac{dy}{dt} + \omega_n^2 y = f(t)"

    render(
        [{"type": "equation", "latex": latex, "mode": "native"}],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    p = document.paragraphs[0]
    omath_paras = p._p.findall(f"{{{OMML_NS}}}oMathPara")
    assert len(omath_paras) == 1
    omath = omath_paras[0].find(f"{{{OMML_NS}}}oMath")
    assert omath is not None

    # Check fraction and superscript
    assert omath.find(f".//{{{OMML_NS}}}f") is not None
    assert omath.find(f".//{{{OMML_NS}}}sSup") is not None


def test_render_matrix_equation_native_docx(
    tmp_path: Path, docx_binding: object
) -> None:
    """Verify state-space matrix equation renders as native OMML in DOCX."""
    output = tmp_path / "matrix_eq.docx"
    latex = (
        r"\begin{bmatrix} \dot{x}_1 \\ \dot{x}_2 \end{bmatrix} = "
        r"\begin{bmatrix} 0 & 1 \\ -\omega_n^2 & -2\zeta\omega_n \end{bmatrix} "
        r"\begin{bmatrix} x_1 \\ x_2 \end{bmatrix} + \begin{bmatrix} 0 \\ 1 \end{bmatrix} u"
    )

    render(
        [{"type": "equation", "latex": latex, "mode": "native"}],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    p = document.paragraphs[0]
    omath_paras = p._p.findall(f"{{{OMML_NS}}}oMathPara")
    assert len(omath_paras) == 1
    omath = omath_paras[0].find(f"{{{OMML_NS}}}oMath")
    assert omath is not None

    # Matrix elements <m:m> and delimiters <m:d>
    matrices = omath.findall(f".//{{{OMML_NS}}}m")
    assert len(matrices) >= 4  # 4 matrices in this equation
    accents = omath.findall(f".//{{{OMML_NS}}}acc")
    assert len(accents) >= 2  # \dot{x}_1 and \dot{x}_2


def test_render_binding_equation_mode_native(
    tmp_path: Path, docx_binding: object
) -> None:
    """Verify setting equation_mode='native' in binding enables native math for all equations."""
    output = tmp_path / "binding_native.docx"
    binding = docx_binding()
    binding["equation_mode"] = "native"

    render(
        [
            {"type": "equation", "latex": r"E = mc^2"},
            {"type": "equation", "latex": r"V_{out} = A_v \cdot V_{in}"},
        ],
        binding,
        output,
    )

    document = Document(str(output))
    assert len(document.paragraphs) == 2
    for p in document.paragraphs:
        assert len(p._p.findall(f"{{{OMML_NS}}}oMathPara")) == 1


@pytest.mark.parametrize(
    "latex",
    [r"\unknowncomplexmacro{x}", "x} + y", r"x & y", r"x \\ y", r"\begin{matrix}a & b"],
)
def test_render_native_preferred_fallback_to_image(
    tmp_path: Path, png_image: Path, docx_binding: object, latex: str
) -> None:
    """Invalid native input falls back to an image in native-preferred mode."""
    output = tmp_path / "fallback.docx"
    # Invalid or unsupported native input carries a pre-rendered image fallback.

    render(
        [
            {
                "type": "equation",
                "latex": latex,
                "mode": "native-preferred",
                "image": str(png_image),
            }
        ],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    # Image fallback should have added an inline shape picture
    assert len(document.inline_shapes) == 1
    assert not document._element.findall(f".//{{{OMML_NS}}}oMath")


def test_render_native_required_preserves_equation_semantics(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "semantics.docx"
    render(
        [
            {
                "type": "equation",
                "latex": r"\begin{bmatrix}f'''_i & \text{ steady  state } \\ c & d\end{bmatrix}",
                "mode": "native-required",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert len(document.inline_shapes) == 0
    math = document.paragraphs[0]._p.find(f"{{{OMML_NS}}}oMathPara/{{{OMML_NS}}}oMath")
    assert math is not None
    rows = math.findall(f".//{{{OMML_NS}}}mr")
    assert [["".join(cell.itertext()) for cell in row] for row in rows] == [
        ["fi′′′", " steady  state "],
        ["c", "d"],
    ]
    script = math.find(f".//{{{OMML_NS}}}sSubSup")
    assert script is not None
    assert "".join(script.find(f"{{{OMML_NS}}}e").itertext()) == "f"
    assert "".join(script.find(f"{{{OMML_NS}}}sub").itertext()) == "i"
    assert "".join(script.find(f"{{{OMML_NS}}}sup").itertext()) == "′′′"
    text_run = next(
        r
        for r in math.findall(f".//{{{OMML_NS}}}r")
        if r.find(f"{{{OMML_NS}}}t").text == " steady  state "
    )
    assert (
        text_run.find(f"{{{OMML_NS}}}t").get(
            "{http://www.w3.org/XML/1998/namespace}space"
        )
        == "preserve"
    )
    assert (
        text_run.find(f"{{{OMML_NS}}}rPr/{{{OMML_NS}}}sty").get(f"{{{OMML_NS}}}val")
        == "p"
    )


def test_render_native_preferred_dynamic_image_fallback(
    tmp_path: Path, docx_binding: object
) -> None:
    """Verify unsupported OMML construct falls back to dynamic matplotlib image rendering."""
    output = tmp_path / "fallback_dynamic.docx"
    # Expression unsupported by our OMML parser but valid in matplotlib mathtext
    latex = r"\boldsymbol{W} + \boldsymbol{X}"

    render(
        [{"type": "equation", "latex": latex, "mode": "native-preferred"}],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    assert len(document.inline_shapes) == 1


@pytest.mark.parametrize(
    "latex",
    [
        r"\unsupportedcomplexmacro{x}",
        "x} + y",
        r"x & y",
        r"x \\ y",
        r"\begin{matrix}a & b",
    ],
)
def test_render_native_required_raises_on_invalid_input(
    tmp_path: Path, docx_binding: object, latex: str
) -> None:
    """Native-required mode rejects invalid input instead of emitting partial math."""
    output = tmp_path / "error.docx"
    with pytest.raises((UnsupportedMacroError, EquationConversionError)):
        render(
            [{"type": "equation", "latex": latex, "mode": "native-required"}],
            docx_binding(),
            output,
        )
