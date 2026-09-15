"""Tests for TOC and pagebreak block rendering."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from presenter.render_docx import render


def test_toc_block_inserts_native_toc_field(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render([{"type": "toc"}], docx_binding(), output)
    document = Document(str(output))
    field_paragraphs = [
        p
        for p in document.paragraphs
        if p._p.findall(".//" + qn("w:fldChar"))
        or p._p.findall(".//" + qn("w:instrText"))
    ]
    assert len(field_paragraphs) == 1
    xml = field_paragraphs[0]._p.xml
    assert "begin" in xml
    assert "separate" in xml
    assert "end" in xml
    assert 'TOC \\o "1-3"' in xml


def test_pagebreak_block_inserts_native_page_break(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render([{"type": "pagebreak"}], docx_binding(), output)
    document = Document(str(output))
    breaks = [
        br
        for p in document.paragraphs
        for br in p._p.findall(".//" + qn("w:br"))
        if br.get(qn("w:type")) == "page"
    ]
    assert len(breaks) == 1


def test_toc_pagebreak_and_text_render_in_order(
    tmp_path: Path, docx_binding: object, paragraph_text: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {"type": "toc"},
            {"type": "pagebreak"},
            {"type": "text", "style": "heading1", "text": "Next section"},
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    children = list(document.element.body.iterchildren())
    toc_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p") and child.findall(".//" + qn("w:instrText"))
    )
    break_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p")
        and any(
            br.get(qn("w:type")) == "page"
            for br in child.findall(".//" + qn("w:br"))
        )
    )
    heading_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p") and paragraph_text(child) == "Next section"
    )
    assert toc_pos < break_pos < heading_pos
