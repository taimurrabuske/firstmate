"""Tests for table-block rendering."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from presenter.render_docx import render


def test_grid_table_with_headers_rows_and_caption(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "caption": "PVT operating points",
                "headers": ["Pressure", "Volume", "Temperature"],
                "rows": [[1.2, 3, 300], [4.5, 6, 600]],
                "style": "grid",
            }
        ],
        docx_binding(),
        output,
    )

    document = Document(str(output))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.style.name == "Table Grid"
    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert [cell.text for cell in table.rows[0].cells] == [
        "Pressure",
        "Volume",
        "Temperature",
    ]
    header_runs = table.rows[0].cells[0].paragraphs[0].runs
    assert header_runs and header_runs[0].bold is True
    assert table.cell(1, 0).text == "1.2"
    assert table.cell(2, 2).text == "600"

    captions = [p for p in document.paragraphs if p.text == "PVT operating points"]
    assert len(captions) == 1
    assert captions[0].style.name in {"Caption", "Normal"}


def test_table_caption_is_placed_above_the_table(
    tmp_path: Path, docx_binding: object, paragraph_text: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "caption": "Spec table",
                "headers": ["A"],
                "rows": [[1]],
                "style": "plain",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    children = list(document.element.body.iterchildren())
    caption_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p") and paragraph_text(child) == "Spec table"
    )
    table_pos = next(
        i for i, child in enumerate(children) if child.tag == qn("w:tbl")
    )
    assert caption_pos < table_pos


def test_striped_style_applies_light_shading(tmp_path: Path, docx_binding: object) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "headers": ["K"],
                "rows": [["v"]],
                "style": "striped",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert document.tables[0].style.name == "Light Shading Accent 1"


def test_plain_style_leaves_default_table_styling(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "headers": ["K"],
                "rows": [["v"]],
                "style": "plain",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert document.tables[0].style.name != "Table Grid"


def test_rows_without_headers(tmp_path: Path, docx_binding: object) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "caption": None,
                "headers": [],
                "rows": [["a", "b"], ["c", "d"]],
                "style": "grid",
            }
        ],
        docx_binding(),
        output,
    )
    table = Document(str(output)).tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.cell(0, 1).text == "b"


def test_ragged_rows_are_padded_to_header_width(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "table",
                "headers": ["X", "Y", "Z"],
                "rows": [["1", "2"]],
                "style": "grid",
            }
        ],
        docx_binding(),
        output,
    )
    table = Document(str(output)).tables[0]
    assert table.cell(1, 2).text == ""


def test_row_wider_than_headers_raises_naming_row_index(
    tmp_path: Path, docx_binding: object
) -> None:
    with pytest.raises(ValueError, match=r"row 1 has 3 cells"):
        render(
            [
                {
                    "type": "table",
                    "headers": ["X", "Y"],
                    "rows": [["1", "2"], ["a", "b", "c"]],
                    "style": "grid",
                }
            ],
            docx_binding(),
            tmp_path / "x.docx",
        )
