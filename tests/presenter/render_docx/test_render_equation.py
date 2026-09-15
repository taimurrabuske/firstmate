"""Tests for equation block rendering."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from presenter.render_docx import render


def test_equation_without_image_falls_back_to_monospace_latex(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "equation", "latex": "V_{out} = \\frac{R_f}{R_i}", "font_size_pt": None}],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "V_{out} = \\frac{R_f}{R_i}"
    assert len(paragraph.runs) == 1
    assert paragraph.runs[0].font.name == "Courier New"


def test_equation_font_size_is_honored(tmp_path: Path, docx_binding: object) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "equation", "latex": "E = mc^2", "font_size_pt": 14.0}],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    run = document.paragraphs[-1].runs[0]
    assert run.font.size == Pt(14.0)
    assert run.font.name == "Courier New"


def test_equation_with_prerendered_image_inserts_picture(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "equation",
                "latex": "ignored when image present",
                "font_size_pt": None,
                "image": str(png_image),
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert len(document.inline_shapes) == 1


def test_equation_image_resolves_artifact_reference(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "equation",
                "latex": "integral form",
                "font_size_pt": None,
                "image": "artifact:eq-01",
            }
        ],
        docx_binding(artifacts={"artifact:eq-01": str(png_image)}),
        output,
    )
    assert len(Document(str(output)).inline_shapes) == 1
