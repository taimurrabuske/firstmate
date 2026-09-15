"""Tests for text-block rendering and style resolution."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from presenter.render_docx import render


def _read(path: Path) -> Document:
    return Document(str(path))


def test_body_text_defaults_to_normal_style(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    result = render(
        [{"type": "text", "style": "body", "text": "Plain body line"}],
        docx_binding(),
        output,
    )
    assert result["blocks_rendered"] == 1
    assert result["format"] == "docx"
    assert result["template"] is None
    assert result["output_path"] == str(output)

    document = _read(output)
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "Plain body line"
    assert paragraph.style.name == "Normal"


def test_heading_styles_resolve_to_builtin_headings(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {"type": "text", "style": "heading1", "text": "Overview"},
            {"type": "text", "style": "heading2", "text": "Electrical specs"},
        ],
        docx_binding(),
        output,
    )
    document = _read(output)
    styles = {p.text: p.style.name for p in document.paragraphs}
    assert styles["Overview"] == "Heading 1"
    assert styles["Electrical specs"] == "Heading 2"


def test_caption_style_defaults_to_caption(tmp_path: Path, docx_binding: object) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "text", "style": "caption", "text": "Fig. 1 preamble"}],
        docx_binding(),
        output,
    )
    document = _read(output)
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "Fig. 1 preamble"
    assert paragraph.style.name in {"Caption", "Normal"}


def test_binding_styles_map_overrides_builtin_names(
    tmp_path: Path,
    make_template_docx: object,
    docx_binding: object,
) -> None:
    template = make_template_docx(custom_styles=("Corporate Body",))
    output = tmp_path / "out.docx"
    render(
        [{"type": "text", "style": "body", "text": "Branded paragraph"}],
        docx_binding(template=template, styles={"body": "Corporate Body"}),
        output,
    )
    document = _read(output)
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "Branded paragraph"
    assert paragraph.style.name == "Corporate Body"


def test_missing_custom_style_falls_back_to_default(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "text", "style": "heading1", "text": "Fallback heading"}],
        docx_binding(styles={"heading1": "NoSuchStyle"}),
        output,
    )
    document = _read(output)
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "Fallback heading"
    assert paragraph.style.name == "Heading 1"


def test_unknown_block_type_raises(tmp_path: Path, docx_binding: object) -> None:
    with pytest.raises(ValueError, match="unsupported block type"):
        render([{"type": "hologram", "text": "?!"}], docx_binding(), tmp_path / "x.docx")


def test_empty_block_list_produces_document(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "empty.docx"
    result = render([], docx_binding(), output)
    assert result["blocks_rendered"] == 0
    assert _read(output) is not None
