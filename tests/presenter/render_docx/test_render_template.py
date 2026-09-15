"""Tests for template binding behavior of the docx engine."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from presenter.render_docx import render


def test_template_binding_starts_from_template_file(
    tmp_path: Path,
    make_template_docx: object,
    docx_binding: object,
) -> None:
    template = make_template_docx(marker="Corporate Master")
    output = tmp_path / "out.docx"
    result = render(
        [{"type": "text", "style": "heading2", "text": "Appendix"}],
        docx_binding(template=template),
        output,
    )
    assert result["template"] == str(template)

    document = Document(str(output))
    texts = [p.text for p in document.paragraphs]
    assert "Corporate Master" in texts  # template content preserved
    assert "Appendix" in texts  # block appended after template content
    heading = next(p for p in document.paragraphs if p.text == "Appendix")
    assert heading.style.name == "Heading 2"


def test_template_styles_apply_to_rendered_blocks(
    tmp_path: Path,
    make_template_docx: object,
    docx_binding: object,
) -> None:
    template = make_template_docx(custom_styles=("Corporate Body",))
    output = tmp_path / "out.docx"
    render(
        [{"type": "text", "style": "body", "text": "Styled from template"}],
        docx_binding(template=template, styles={"body": "Corporate Body"}),
        output,
    )
    document = Document(str(output))
    paragraph = next(
        p for p in document.paragraphs if p.text == "Styled from template"
    )
    assert paragraph.style.name == "Corporate Body"


def test_non_docx_binding_format_is_rejected(
    tmp_path: Path, docx_binding: object
) -> None:
    binding = docx_binding()
    binding["format"] = "pptx"
    with pytest.raises(ValueError, match="format 'docx', got 'pptx'"):
        render([{"type": "toc"}], binding, tmp_path / "x.docx")


def test_empty_template_field_renders_blank_document(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    result = render(
        [{"type": "text", "style": "body", "text": "hello"}],
        docx_binding(template=None),
        output,
    )
    assert result["template"] is None
    document = Document(str(output))
    assert document.paragraphs[-1].text == "hello"


def test_none_binding_renders_blank_document(tmp_path: Path) -> None:
    output = tmp_path / "out.docx"
    result = render([{"type": "pagebreak"}], None, output)
    assert result["blocks_rendered"] == 1
    assert Document(str(output)) is not None


def test_output_is_written_into_nested_directory(
    tmp_path: Path, docx_binding: object
) -> None:
    output = tmp_path / "reports" / "deep" / "datasheet.docx"
    result = render([], docx_binding(), output)
    assert result["output_path"] == str(output)
    assert output.is_file()
