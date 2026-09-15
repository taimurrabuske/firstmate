"""Tests for image and plot block rendering."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from presenter.render_docx import render


def _captions(document: Document) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name == "Caption"]


def test_image_block_inserts_picture_with_width_and_caption(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "image",
                "source": str(png_image),
                "width_in": 2.5,
                "caption": "Board schematic",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width == Inches(2.5)
    assert "Board schematic" in _captions(document)


def test_image_block_without_width_still_inserts(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "image", "source": str(png_image), "width_in": None, "caption": None}],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert len(document.inline_shapes) == 1


def test_plot_block_behaves_like_image(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "plot",
                "source": str(png_image),
                "width_in": 4.0,
                "caption": "Load transient waveform",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width == Inches(4.0)
    assert "Load transient waveform" in _captions(document)


def test_artifact_reference_resolves_through_binding(
    tmp_path: Path, png_image: Path, docx_binding: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [{"type": "plot", "source": "artifact:wave-01", "width_in": 3.0, "caption": None}],
        docx_binding(artifacts={"artifact:wave-01": str(png_image)}),
        output,
    )
    assert len(Document(str(output)).inline_shapes) == 1


def test_unknown_artifact_reference_raises(
    tmp_path: Path, docx_binding: object
) -> None:
    with pytest.raises(FileNotFoundError, match="artifact:missing"):
        render(
            [{"type": "image", "source": "artifact:missing", "width_in": None, "caption": None}],
            docx_binding(),
            tmp_path / "x.docx",
        )


def test_mapped_artifact_with_missing_target_names_mapped_path(
    tmp_path: Path, docx_binding: object
) -> None:
    target = tmp_path / "missing-target.png"
    with pytest.raises(FileNotFoundError, match="missing-target.png"):
        render(
            [
                {
                    "type": "image",
                    "source": "artifact:wave-01",
                    "width_in": None,
                    "caption": None,
                }
            ],
            docx_binding(artifacts={"artifact:wave-01": str(target)}),
            tmp_path / "x.docx",
        )


def test_picture_caption_follows_the_picture_paragraph(
    tmp_path: Path, png_image: Path, docx_binding: object, paragraph_text: object
) -> None:
    output = tmp_path / "out.docx"
    render(
        [
            {
                "type": "image",
                "source": str(png_image),
                "width_in": None,
                "caption": "Fig. 2",
            }
        ],
        docx_binding(),
        output,
    )
    document = Document(str(output))
    children = list(document.element.body.iterchildren())
    picture_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p") and child.findall(".//" + qn("w:drawing"))
    )
    caption_pos = next(
        i
        for i, child in enumerate(children)
        if child.tag == qn("w:p") and paragraph_text(child) == "Fig. 2"
    )
    assert caption_pos == picture_pos + 1


def test_image_block_without_source_raises(tmp_path: Path, docx_binding: object) -> None:
    with pytest.raises(ValueError, match="requires a 'source' key"):
        render(
            [{"type": "image", "width_in": None, "caption": None}],
            docx_binding(),
            tmp_path / "x.docx",
        )
