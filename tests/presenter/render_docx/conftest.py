"""Shared fixtures for presenter.render_docx tests.

Builds every .docx and image fixture in-test with python-docx and Pillow;
no binary fixtures are committed.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Callable

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from PIL import Image

REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


@pytest.fixture()
def paragraph_text() -> Callable[..., str]:
    """Extract the concatenated w:t text of a body element via python-docx XML."""

    def _text(element: object) -> str:
        return "".join(
            node.text or "" for node in element.iter(qn("w:t"))  # type: ignore[attr-defined]
        )

    return _text


@pytest.fixture()
def png_image(tmp_path: Path) -> Path:
    """A tiny real PNG so python-docx can embed it."""
    path = tmp_path / "figure.png"
    Image.new("RGB", (32, 16), color=(180, 40, 40)).save(path, format="PNG")
    return path


@pytest.fixture()
def make_template_docx(tmp_path: Path) -> Callable[..., Path]:
    """Build a minimal corporate .docx template carrying a marker and custom styles."""

    def _make(
        name: str = "template.docx",
        marker: str = "Corporate Master",
        custom_styles: tuple[str, ...] = (),
    ) -> Path:
        document = Document()
        document.add_heading(marker, level=1)
        document.add_paragraph("Confidential drafting note", style="Intense Quote")
        for style_name in custom_styles:
            document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        path = tmp_path / name
        document.save(str(path))
        return path

    return _make


@pytest.fixture()
def docx_binding() -> Callable[..., dict[str, Any]]:
    """Build a TemplateBinding dict per the shared contract."""

    def _make(
        template: Path | None = None,
        styles: dict[str, str] | None = None,
        artifacts: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        binding: dict[str, Any] = {
            "format": "docx",
            "template": str(template) if template else "",
            "name": "Test Binding",
            "layouts": {},
            "placeholders": {},
            "styles": dict(styles or {}),
        }
        if artifacts:
            binding["artifacts"] = dict(artifacts)
        return binding

    return _make
