"""Tests for presenter._ooxml.svg and OOXML constants."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc import serialized
from pptx.oxml import parse_xml

from presenter._ooxml.constants import (
    CONTENT_TYPE_PNG,
    CONTENT_TYPE_SVG,
    NS_A,
    NS_ASVG,
    NS_R,
    RELATIONSHIP_TYPE_IMAGE,
    SVG_EXTENSION_URI,
)
from presenter._ooxml.svg import (
    add_svg_picture,
    attach_svg_blip_extension,
    ensure_svg_content_type,
    get_or_add_svg_part,
    is_svg_asset,
    relate_svg_part,
)

_SAMPLE_SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
    b'<rect width="200" height="100" fill="blue"/></svg>'
)


def test_ooxml_constants() -> None:
    """Verify standard Microsoft extension URI and namespaces."""
    assert SVG_EXTENSION_URI == "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
    assert CONTENT_TYPE_SVG == "image/svg+xml"
    assert CONTENT_TYPE_PNG == "image/png"
    assert RELATIONSHIP_TYPE_IMAGE == (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    )


def test_ensure_svg_content_type() -> None:
    """Verify SVG content-type mapping is registered in python-pptx defaults."""
    ensure_svg_content_type()
    assert ("svg", CONTENT_TYPE_SVG) in serialized.default_content_types


def test_is_svg_asset(tmp_path: Path) -> None:
    """Verify detection of SVG assets by path, filename, and bytes content."""
    svg_file = tmp_path / "diagram.svg"
    svg_file.write_bytes(_SAMPLE_SVG)
    png_file = tmp_path / "photo.png"
    png_file.write_bytes(b"\x89PNG\r\n\x1a\n...")

    assert is_svg_asset("diagram.svg") is True
    assert is_svg_asset(svg_file) is True
    assert is_svg_asset(_SAMPLE_SVG) is True
    assert is_svg_asset("photo.png") is False
    assert is_svg_asset(png_file) is False
    assert is_svg_asset(b"\x89PNG...") is False


def test_get_or_add_svg_part_deduplication() -> None:
    """Verify SVG parts are deduplicated across a presentation package by sha256."""
    prs = Presentation()
    package = prs.part.package

    part1 = get_or_add_svg_part(package, _SAMPLE_SVG)
    part2 = get_or_add_svg_part(package, _SAMPLE_SVG)

    assert part1 is part2
    assert part1.content_type == CONTENT_TYPE_SVG
    assert part1.partname.endswith(".svg")

    different_svg = b'<svg xmlns="http://www.w3.org/2000/svg"><circle r="5"/></svg>'
    part3 = get_or_add_svg_part(package, different_svg)
    assert part3 is not part1
    assert part3.partname != part1.partname


def test_relate_svg_part() -> None:
    """Verify relating an SVG part to a slide produces an image relationship."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    svg_part = get_or_add_svg_part(prs.part.package, _SAMPLE_SVG)

    r_id = relate_svg_part(slide.part, svg_part)
    assert r_id.startswith("rId")

    # Re-relating the same part to the same slide returns the same relationship ID
    r_id2 = relate_svg_part(slide.part, svg_part)
    assert r_id2 == r_id


def test_attach_svg_blip_extension_new_ext_lst() -> None:
    """Verify attaching asvg:svgBlip when a:blip has no prior a:extLst."""
    blip = parse_xml(f'<a:blip xmlns:a="{NS_A}" xmlns:r="{NS_R}" r:embed="rId2"/>')
    attach_svg_blip_extension(blip, "rId5")

    ext_lst = blip.find(f"{{{NS_A}}}extLst")
    assert ext_lst is not None
    ext = ext_lst.find(f"./{{{NS_A}}}ext[@uri='{SVG_EXTENSION_URI}']")
    assert ext is not None
    svg_blip = ext.find(f"{{{NS_ASVG}}}svgBlip")
    assert svg_blip is not None
    assert svg_blip.get(f"{{{NS_R}}}embed") == "rId5"


def test_attach_svg_blip_extension_existing_ext_lst() -> None:
    """Verify attaching asvg:svgBlip when an existing a:extLst has other extensions."""
    blip = parse_xml(
        f'<a:blip xmlns:a="{NS_A}" xmlns:r="{NS_R}" r:embed="rId2">'
        f'  <a:extLst><a:ext uri="{{OTHER-EXTENSION-URI}}"/></a:extLst>'
        f'</a:blip>'
    )
    attach_svg_blip_extension(blip, "rId8")

    ext = blip.find(f".//{{{NS_A}}}ext[@uri='{SVG_EXTENSION_URI}']")
    assert ext is not None
    svg_blip = ext.find(f"{{{NS_ASVG}}}svgBlip")
    assert svg_blip is not None
    assert svg_blip.get(f"{{{NS_R}}}embed") == "rId8"


def test_attach_svg_blip_extension_update_existing() -> None:
    """Verify updating an existing asvg:svgBlip relationship ID."""
    blip = parse_xml(
        f'<a:blip xmlns:a="{NS_A}" xmlns:r="{NS_R}" r:embed="rId2">'
        f'  <a:extLst>'
        f'    <a:ext uri="{SVG_EXTENSION_URI}">'
        f'      <asvg:svgBlip xmlns:asvg="{NS_ASVG}" xmlns:r="{NS_R}" r:embed="rIdOld"/>'
        f'    </a:ext>'
        f'  </a:extLst>'
        f'</a:blip>'
    )
    attach_svg_blip_extension(blip, "rIdNew")

    svg_blip = blip.find(f".//{{{NS_ASVG}}}svgBlip")
    assert svg_blip is not None
    assert svg_blip.get(f"{{{NS_R}}}embed") == "rIdNew"


def test_add_svg_picture_end_to_end(tmp_path: Path) -> None:
    """Verify high-level add_svg_picture creates a valid PPTX with dual relationships."""
    from PIL import Image

    # Create dummy PNG fallback
    png_bytes = io.BytesIO()
    Image.new("RGB", (200, 100), color="blue").save(png_bytes, format="PNG")
    png_data = png_bytes.getvalue()

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    pic = add_svg_picture(
        slide,
        svg_source=_SAMPLE_SVG,
        png_fallback=png_data,
        left=100,
        top=100,
        alt_text="A vector diagram",
    )

    assert pic is not None
    blip = pic._element.blipFill.blip
    svg_blip = blip.find(f".//{{{NS_ASVG}}}svgBlip")
    assert svg_blip is not None
    svg_r_id = svg_blip.get(f"{{{NS_R}}}embed")
    assert svg_r_id is not None
    assert svg_r_id != blip.get(f"{{{NS_R}}}embed")

    # Save to buffer and verify zip package
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    with zipfile.ZipFile(buf) as z:
        namelist = z.namelist()
        svg_parts = [n for n in namelist if n.endswith(".svg")]
        png_parts = [n for n in namelist if n.endswith(".png")]
        assert len(svg_parts) == 1
        assert len(png_parts) == 1
        assert z.read(svg_parts[0]) == _SAMPLE_SVG

        # Check [Content_Types].xml
        ct = z.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="svg" ContentType="image/svg+xml"' in ct
