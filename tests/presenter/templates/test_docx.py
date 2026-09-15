"""Tests for presenter.templates.docx, built against an in-test docx fixture."""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from presenter.templates.binding import validate_binding
from presenter.templates.docx import inventory_docx_template


def _add_content_control(document, tag, alias, text):
    """Insert a plain-text structured document tag (content control) into
    the document body, immediately before the section properties element
    so the resulting document stays well-formed.
    """
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")

    tag_el = OxmlElement("w:tag")
    tag_el.set(qn("w:val"), tag)
    alias_el = OxmlElement("w:alias")
    alias_el.set(qn("w:val"), alias)
    sdt_pr.append(tag_el)
    sdt_pr.append(alias_el)
    sdt_pr.append(OxmlElement("w:text"))

    sdt_content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    run_text = OxmlElement("w:t")
    run_text.text = text
    run.append(run_text)
    paragraph.append(run)
    sdt_content.append(paragraph)

    sdt.append(sdt_pr)
    sdt.append(sdt_content)

    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(sdt)
    else:
        body.append(sdt)


@pytest.fixture
def datasheet_template_path(tmp_path):
    document = Document()
    document.add_paragraph("Cover page", style="Title")
    _add_content_control(document, "part_number", "Part Number", "XYZ-100")
    _add_content_control(document, "revision", "Revision", "A")

    path = tmp_path / "corp_datasheet.docx"
    document.save(str(path))
    return path


def test_inventory_docx_template_reports_format_and_identity(datasheet_template_path):
    binding = inventory_docx_template(datasheet_template_path)

    assert binding["format"] == "docx"
    assert binding["template"] == str(datasheet_template_path)
    assert binding["name"] == "corp_datasheet"


def test_inventory_docx_template_detects_content_controls(datasheet_template_path):
    binding = inventory_docx_template(datasheet_template_path)

    controls = binding["placeholders"]["content_controls"]
    tags = {control["tag"] for control in controls}
    assert tags == {"part_number", "revision"}
    for control in controls:
        assert control["type"] == "text"


def test_inventory_docx_template_maps_placeholder_roles_by_tag(
    datasheet_template_path,
):
    binding = inventory_docx_template(datasheet_template_path)

    default_layout = binding["layouts"]["default"]
    assert default_layout["part_number"] == "part_number"
    assert default_layout["revision"] == "revision"


def test_inventory_docx_template_with_no_content_controls_yields_empty_layout(
    tmp_path,
):
    document = Document()
    document.add_paragraph("Plain document, no fields.")
    path = tmp_path / "plain.docx"
    document.save(str(path))

    binding = inventory_docx_template(path)

    assert binding["layouts"]["default"] == {}
    assert binding["placeholders"]["content_controls"] == []


def test_inventory_docx_template_lists_builtin_styles(datasheet_template_path):
    binding = inventory_docx_template(datasheet_template_path)

    assert "Title" in binding["styles"]["paragraph"]
    assert "Normal" in binding["styles"]["paragraph"]
    assert "Default Paragraph Font" in binding["styles"]["character"]


def test_inventory_docx_template_reports_section_geometry(datasheet_template_path):
    binding = inventory_docx_template(datasheet_template_path)

    sections = binding["styles"]["sections"]
    assert len(sections) == 1
    assert sections[0]["page_width_emu"] > 0
    assert sections[0]["page_height_emu"] > 0
    assert sections[0]["orientation"] == "PORTRAIT"


def test_inventory_docx_template_produces_a_valid_binding(datasheet_template_path):
    binding = inventory_docx_template(datasheet_template_path)
    validate_binding(binding)
