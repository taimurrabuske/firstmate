"""Tests for presenter.templates.pptx, built against an in-test pptx fixture."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart

from presenter.templates.binding import validate_binding
from presenter.templates.pptx import inventory_pptx_template


_OFFICE_NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
)

_DARK_MASTER_XML = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<p:sldMaster {_OFFICE_NS}>
  <p:cSld name="Corporate Dark Master">
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr/>
    </p:spTree>
  </p:cSld>
  <p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1"
            accent2="accent2" accent3="accent3" accent4="accent4"
            accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>
  <p:sldLayoutIdLst/>
</p:sldMaster>"""

_DARK_LAYOUT_XML = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<p:sldLayout {_OFFICE_NS} type="title">
  <p:cSld name="Dark Title Slide">
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr/>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="2" name="Dark Title"/>
          <p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
          <p:nvPr><p:ph type="title"/></p:nvPr>
        </p:nvSpPr>
        <p:spPr/>
        <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
  <p:clrOvr><a:masterClrMapping/></p:clrOvr>
</p:sldLayout>"""


def _save_two_master_template(path: Path) -> None:
    """Build a real two-master .pptx: the stock Office master plus a second
    "Corporate Dark" master with its own layout and its own theme accent."""
    prs = Presentation()
    prs.slide_masters[0].name = "Standard Master"

    package = prs.part.package
    stock_theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
    dark_theme = Part(
        PackURI("/ppt/theme/theme2.xml"),
        stock_theme.content_type,
        package,
        stock_theme.blob.replace(
            b'<a:accent1><a:srgbClr val="4F81BD"/>',
            b'<a:accent1><a:srgbClr val="00B0F0"/>',
        ),
    )
    dark_layout = SlideLayoutPart.load(
        PackURI("/ppt/slideLayouts/slideLayout12.xml"),
        CONTENT_TYPE.PML_SLIDE_LAYOUT,
        package,
        _DARK_LAYOUT_XML.encode("utf-8"),
    )
    dark_master = SlideMasterPart.load(
        PackURI("/ppt/slideMasters/slideMaster2.xml"),
        CONTENT_TYPE.PML_SLIDE_MASTER,
        package,
        _DARK_MASTER_XML.encode("utf-8"),
    )

    layout_rid = dark_master.relate_to(dark_layout, RT.SLIDE_LAYOUT)
    dark_master.relate_to(dark_theme, RT.THEME)
    layout_id_lst = dark_master.slide_master.element.get_or_add_sldLayoutIdLst()
    layout_id_lst.append(
        layout_id_lst.makeelement(
            qn("p:sldLayoutId"), {"id": "2147484000", qn("r:id"): layout_rid}
        )
    )

    master_rid = prs.part.relate_to(dark_master, RT.SLIDE_MASTER)
    master_id_lst = prs.element.get_or_add_sldMasterIdLst()
    master_id_lst.append(
        master_id_lst.makeelement(
            qn("p:sldMasterId"), {"id": "2147484001", qn("r:id"): master_rid}
        )
    )

    prs.save(str(path))


@pytest.fixture
def master_template_path(tmp_path):
    """A default python-pptx template, which already ships a full layout set
    (Title Slide, Title and Content, etc.) and a standard Office theme."""
    presentation = Presentation()
    path = tmp_path / "corp_master.pptx"
    presentation.save(str(path))
    return path


def test_inventory_pptx_template_reports_format_and_identity(master_template_path):
    binding = inventory_pptx_template(master_template_path)

    assert binding["format"] == "pptx"
    assert binding["template"] == str(master_template_path)
    assert binding["name"] == "corp_master"


def test_inventory_pptx_template_lists_every_layout(master_template_path):
    binding = inventory_pptx_template(master_template_path)
    presentation = Presentation(str(master_template_path))

    expected_layout_names = {layout.name for layout in presentation.slide_layouts}
    assert set(binding["layouts"]) == expected_layout_names
    assert set(binding["placeholders"]) == expected_layout_names


def test_inventory_pptx_template_maps_title_slide_placeholder_roles(
    master_template_path,
):
    binding = inventory_pptx_template(master_template_path)

    title_slide_roles = binding["layouts"]["Title Slide"]
    assert title_slide_roles["Title 1"] == "title"
    assert title_slide_roles["Subtitle 2"] == "subtitle"


def test_inventory_pptx_template_records_placeholder_type_and_idx(
    master_template_path,
):
    binding = inventory_pptx_template(master_template_path)

    title_slide_inventory = binding["placeholders"]["Title Slide"]
    by_name = {entry["name"]: entry for entry in title_slide_inventory}

    assert by_name["Title 1"]["type"] == "CENTER_TITLE"
    assert by_name["Title 1"]["idx"] == 0
    assert by_name["Subtitle 2"]["type"] == "SUBTITLE"
    assert by_name["Subtitle 2"]["idx"] == 1


def test_inventory_pptx_template_reports_theme_fonts_and_colors(
    master_template_path,
):
    binding = inventory_pptx_template(master_template_path)

    theme = binding["styles"]["theme"]
    assert theme["fonts"]["major"]
    assert theme["fonts"]["minor"]
    # The default python-pptx template uses the stock Office palette.
    assert theme["colors"]["accent1"].upper() == "4F81BD"


def test_inventory_pptx_template_resolves_system_colors_to_concrete_hex(
    master_template_path,
):
    binding = inventory_pptx_template(master_template_path)

    colors = binding["styles"]["theme"]["colors"]
    assert colors["dk1"] == "000000"
    assert colors["lt1"] == "FFFFFF"
    assert set(colors["dk1"]) <= set("0123456789ABCDEFabcdef")
    assert set(colors["lt1"]) <= set("0123456789ABCDEFabcdef")


def test_inventory_pptx_template_reports_slide_size(master_template_path):
    binding = inventory_pptx_template(master_template_path)

    slide_size = binding["styles"]["slide_size"]
    assert slide_size["width_emu"] > 0
    assert slide_size["height_emu"] > 0


def test_inventory_pptx_template_produces_a_valid_binding(master_template_path):
    binding = inventory_pptx_template(master_template_path)
    validate_binding(binding)


def test_inventory_pptx_template_reports_single_master_in_masters_array(
    master_template_path,
):
    binding = inventory_pptx_template(master_template_path)

    assert len(binding["masters"]) == 1
    assert binding["masters"][0]["layouts"] == binding["layouts"]
    assert binding["masters"][0]["theme"] == binding["styles"]["theme"]


def test_inventory_pptx_template_enumerates_every_slide_master(tmp_path):
    path = tmp_path / "multi_master.pptx"
    _save_two_master_template(path)

    binding = inventory_pptx_template(path)
    presentation = Presentation(str(path))
    masters = binding["masters"]

    assert len(masters) == len(presentation.slide_masters) == 2
    assert [master["name"] for master in masters] == [
        "Standard Master",
        "Corporate Dark Master",
    ]
    assert masters[0]["layouts"]["Title Slide"]["Title 1"] == "title"
    assert masters[1]["layouts"]["Dark Title Slide"]["Dark Title"] == "title"
    assert masters[0]["theme"]["colors"]["accent1"] == "4F81BD"
    assert masters[1]["theme"]["colors"]["accent1"] == "00B0F0"
    assert masters[1]["theme"]["fonts"]["major"]


def test_inventory_pptx_template_keeps_first_master_as_top_level_view(tmp_path):
    path = tmp_path / "multi_master.pptx"
    _save_two_master_template(path)

    binding = inventory_pptx_template(path)

    assert "Title Slide" in binding["layouts"]
    assert "Dark Title Slide" not in binding["layouts"]
    assert "Dark Title Slide" not in binding["placeholders"]
    assert binding["styles"]["theme"]["colors"]["accent1"] == "4F81BD"
    validate_binding(binding)
