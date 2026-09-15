#!/usr/bin/env python3
"""Render a real .docx datasheet through the full presenter pipeline.

Assembles a spec from every content-block type (including a table from
``presenter.blocks.tables``, a plot from ``presenter.blocks.plot``, and an
equation), binds it to a fixture corporate .docx template inventoried at
runtime, resolves one image through the artifact source lane, and renders
the result via ``presenter.core.render``.

Every generated file - the fixture template, the artifact image, the
rendered plot and equation pictures, and the final .docx - lives under a
temporary directory printed at the end; nothing is written into the repo.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from PIL import Image

import presenter.render_docx
from presenter.artifacts import FilesystemArtifactSource
from presenter.blocks.equation import build_equation
from presenter.blocks.plot import build_plot
from presenter.blocks.tables import SpecItem, build_spec_table
from presenter.core import DocumentBuilder, render
from presenter.core.blocks import image, pagebreak, text, toc
from presenter.templates.docx import inventory_docx_template


def _build_fixture_template(path: Path) -> None:
    """Write a minimal corporate .docx template to stand in for a real one."""
    document = Document()
    document.add_heading("ACME Corporate Datasheet Template", level=1)
    document.add_paragraph("Confidential - internal engineering use only.")
    document.save(str(path))


def _build_fixture_artifact(artifacts_root: Path, ref: str) -> None:
    """Write a placeholder schematic PNG for the artifact source to resolve."""
    path = artifacts_root / ref
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (320, 180), color=(60, 90, 140)).save(path, format="PNG")


def main() -> Path:
    workdir = Path(tempfile.mkdtemp(prefix="presenter-datasheet-demo-"))

    # The DOCX engine only dispatches through presenter.core.render once
    # registered; importing presenter.render_docx alone never does this.
    presenter.render_docx.register()

    # Bind a fixture corporate template the way the template lane would
    # inventory a real one.
    template_path = workdir / "acme-template.docx"
    _build_fixture_template(template_path)
    binding = inventory_docx_template(template_path)

    # Resolve one image through the artifact source lane rather than a
    # literal filesystem path, then hand its resolved path to the DOCX
    # engine through binding["artifacts"] (its forward-compatible hook for
    # this lane, documented in presenter/render_docx/engine.py).
    artifacts_root = workdir / "artifacts"
    schematic_ref = "schematics/ldo-top.png"
    _build_fixture_artifact(artifacts_root, schematic_ref)
    artifact_source = FilesystemArtifactSource(artifacts_root)
    binding["artifacts"] = {schematic_ref: str(artifact_source.resolve_path(schematic_ref))}

    # Technical content blocks from presenter.blocks.
    plot_block = build_plot(
        x=[0.0, 0.2, 0.4, 0.6, 0.8, 1.0],
        y=[3.30, 3.28, 3.05, 3.10, 3.29, 3.30],
        title="Load step response",
        x_label="Time",
        x_unit="ms",
        y_label="Output voltage",
        y_unit="V",
        caption="Load step 10 mA to 500 mA",
        width_in=6.0,
        output_path=workdir / "load-step.png",
    )
    equation_block = build_equation(
        r"V_{out} = V_{ref} \left(1 + \frac{R_1}{R_2}\right)",
        output_path=workdir / "equation.png",
        font_size_pt=12,
    )
    table_block = build_spec_table(
        [
            SpecItem(
                "Output voltage", symbol="V_OUT",
                min_val=3.234, typ_val=3.3, max_val=3.366, unit="V",
            ),
            SpecItem(
                "Dropout voltage", symbol="V_DO",
                max_val=250, unit="mV", conditions="I_OUT = 500 mA",
            ),
        ],
        caption="Electrical characteristics over PVT unless noted",
    )

    spec = (
        DocumentBuilder("ACME LDO-3V3 Datasheet", author="Analog IC team")
        .slide("Contents").add(toc())
        .slide("Overview")
        .add(text("A 3.3 V, 500 mA low-dropout regulator.", style="body"))
        .add(image(schematic_ref, width_in=5.0, caption="Top-level schematic"))
        .slide("Electrical characteristics")
        .add(table_block)
        .add(equation_block)
        .add(pagebreak())
        .slide("Transient response")
        .add(plot_block)
        .build()
    )

    output_path = workdir / "ldo-3v3-datasheet.docx"
    written = render(spec, binding, output_path)
    print(f"Rendered datasheet: {written}")
    return written


if __name__ == "__main__":
    main()
